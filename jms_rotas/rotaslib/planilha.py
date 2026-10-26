"""Planilha no padrão J&T de um painel (troncal ou secundária): cabeçalho vermelho com português em cima e mandarim
embaixo, aba Resumo com os cartões e os gráficos do Excel, e as abas de dados (pacotes errados, todos os pacotes, viagens).

Gravada em modo "write-only" do openpyxl: aguenta dezenas de milhares de pacotes sem pesar.
"""

from __future__ import annotations

import io
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .analise import contar

VERMELHO = config.VERMELHO_JT.lstrip("#")
_FONTE = "Arial"
_CAB_PT = PatternFill("solid", fgColor=VERMELHO)
_CAB_ZH = PatternFill("solid", fgColor="B8000E")
_BRANCO = Font(name=_FONTE, color="FFFFFF", bold=True, size=10)
_BRANCO_ZH = Font(name=_FONTE, color="FFFFFF", size=9)
_TITULO = Font(name=_FONTE, bold=True, size=16, color=VERMELHO)
_SUB = Font(name=_FONTE, bold=True, size=11, color=VERMELHO)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DATAS = {"tempo_digitalizacao", "saida_planejada", "saida_real", "chegada_real"}
_LARGURAS = {"remessa": 18, "linha": 26, "status": 13, "lote": 17, "base_escaneamento": 14, "proxima_parada": 16,
             "base_destino": 16, "municipio_destino": 20, "tipo_bipagem": 20, "digitalizado_por": 28, "tempo_digitalizacao": 20,
             "turno": 8, "segmento": 14, "destino": 18, "rota": 16, "situacao": 22, "regra": 34, "id_viagem": 18,
             "origem": 14, "destino_viagem": 16, "saida_planejada": 20, "saida_real": 20, "chegada_real": 20, "placa": 10,
             "transportadora": 22, "bipados": 12, "baixados": 10, "errados": 10, "pct_erro": 9, "consulta": 30}


def _cel(ws, valor, fonte=None, preenchimento=None, alinhamento=None, formato=None) -> WriteOnlyCell:
    c = WriteOnlyCell(ws, value=valor)
    if fonte:
        c.font = fonte
    if preenchimento:
        c.fill = preenchimento
    if alinhamento:
        c.alignment = alinhamento
    if formato:
        c.number_format = formato
    return c


def _cabecalho_duplo(ws, chaves: list[str], recuo: int = 0) -> None:
    vazio = [None] * recuo
    ws.append(vazio + [_cel(ws, config.COLUNAS.get(k, (k, k))[0], _BRANCO, _CAB_PT, _CENTRO) for k in chaves])
    ws.append(vazio + [_cel(ws, config.COLUNAS.get(k, (k, k))[1], _BRANCO_ZH, _CAB_ZH, _CENTRO) for k in chaves])


def _aba_dados(wb: Workbook, titulo: str, df: pd.DataFrame, chaves: list[str]) -> None:
    ws = wb.create_sheet(titulo)
    for j, k in enumerate(chaves, start=1):
        ws.column_dimensions[get_column_letter(j)].width = _LARGURAS.get(k, 16)
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(chaves))}{max(2, len(df) + 2)}"
    _cabecalho_duplo(ws, chaves)
    dados = df[chaves].copy()
    # Horários viram data do Excel (filtro por data/hora); o que não for horário fica como texto.
    for k in _DATAS.intersection(chaves):
        convertido = pd.to_datetime(dados[k], format="%Y-%m-%d %H:%M:%S", errors="coerce")
        dados[k] = [d.to_pydatetime() if not pd.isna(d) else (o or None) for d, o in zip(convertido, dados[k])]
    if "pct_erro" in chaves:
        dados["pct_erro"] = dados["pct_erro"] / 100
    i_sit = chaves.index("situacao") if "situacao" in chaves else -1
    i_pct = chaves.index("pct_erro") if "pct_erro" in chaves else -1
    vermelho = Font(name=_FONTE, bold=True, color=config.COR_CRITICO.lstrip("#"))
    for registro in dados.itertuples(index=False, name=None):
        linha = list(registro)
        if i_sit >= 0 and linha[i_sit] == config.ERRADO:
            linha[i_sit] = _cel(ws, config.ERRADO, vermelho)
        if i_pct >= 0:
            linha[i_pct] = _cel(ws, linha[i_pct], formato="0.00%")
        ws.append(linha)


def _grafico_barras(ws, titulo: str, col_cat: int, col_val: int, primeira: int, ultima: int, ancora: str,
                    cores: list[str] | None = None, horizontal: bool = True) -> None:
    g = BarChart()
    g.type = "bar" if horizontal else "col"
    g.title = titulo
    g.style = 10
    g.width, g.height = 17.0, max(7.5, min(20.0, 1.4 + 0.55 * (ultima - primeira + 1))) if horizontal else 7.5
    g.add_data(Reference(ws, min_col=col_val, min_row=primeira, max_row=ultima), titles_from_data=False)
    g.set_categories(Reference(ws, min_col=col_cat, min_row=primeira, max_row=ultima))
    s = g.series[0]
    s.graphicalProperties.solidFill = VERMELHO
    s.graphicalProperties.line.solidFill = VERMELHO
    for i, cor in enumerate(cores or []):
        p = DataPoint(idx=i)
        p.graphicalProperties.solidFill = cor.lstrip("#")
        p.graphicalProperties.line.solidFill = cor.lstrip("#")
        s.dPt.append(p)
    g.dataLabels = DataLabelList()
    g.dataLabels.showVal = True
    for campo in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(g.dataLabels, campo, False)
    g.legend = None
    g.gapWidth = 60
    g.x_axis.delete = False
    g.y_axis.delete = False
    if horizontal:
        g.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(g, ancora)


def _grafico_rosca(ws, titulo: str, primeira: int, ultima: int, ancora: str, cores: list[str]) -> None:
    g = DoughnutChart(holeSize=55)
    g.title = titulo
    g.width, g.height = 12.0, 7.5
    g.add_data(Reference(ws, min_col=2, min_row=primeira, max_row=ultima), titles_from_data=False)
    g.set_categories(Reference(ws, min_col=1, min_row=primeira, max_row=ultima))
    s = g.series[0]
    for i, cor in enumerate(cores):
        p = DataPoint(idx=i)
        p.graphicalProperties.solidFill = cor.lstrip("#")
        s.dPt.append(p)
    g.dataLabels = DataLabelList()
    g.dataLabels.showVal = True
    g.dataLabels.showPercent = True
    for campo in ("showSerName", "showCatName", "showLegendKey"):
        setattr(g.dataLabels, campo, False)
    ws.add_chart(g, ancora)


def gerar_planilha(tipo: str, pacotes: pd.DataFrame, viagens: pd.DataFrame, numeros: dict[str, Any], info: dict[str, Any]) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Resumo 汇总")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 11
    # Impressão/PDF: a largura toda (tabelas e gráficos) numa página.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    linha = 0

    def escrever(*valores) -> None:
        nonlocal linha
        ws.append(list(valores))
        linha += 1

    escrever(_cel(ws, f"J&T Express · Expedição errada · {config.NOMES[tipo]} · {config.NOMES_ZH[tipo]}错发", _TITULO))
    escrever(f"Dia {info.get('dia', '')} · base {info.get('base', config.BASE_NOME)} · atualizado em {info.get('atualizado_em', '')} · "
             f"Turnos: " + " · ".join(f"{t} {h}" for t, h in config.HORARIO_TURNOS.items()))
    escrever()
    escrever(_cel(ws, "Resumo", _SUB))
    ws.append([_cel(ws, "Indicador", _BRANCO, _CAB_PT, _CENTRO), _cel(ws, "Valor", _BRANCO, _CAB_PT, _CENTRO)])
    ws.append([_cel(ws, "指标", _BRANCO_ZH, _CAB_ZH, _CENTRO), _cel(ws, "值", _BRANCO_ZH, _CAB_ZH, _CENTRO)])
    linha += 2
    n = numeros
    pares = [
        ("Viagens concluídas · 已完成", n["viagens_concluidas"]), ("Viagens em trânsito · 在途", n["viagens_transito"]),
        ("Pacotes conferidos · 已核对", n["pacotes"]), ("Pacotes errados · 错发", n["errados"]),
        ("% errado · 错发率", n["pct_erro"] / 100),
        *[(f"Turno {t} ({config.HORARIO_TURNOS[t]}): errados · 错发", n["turnos"][t]["errados"]) for t in config.TURNOS],
        *[(f"Turno {t}: participação nos erros · 占比", n["turnos"][t]["participacao"] / 100) for t in config.TURNOS],
        ("Base de escaneamento mais afetada · 扫描网点", f"{n['base_escaneamento'][0]} ({n['base_escaneamento'][1]})"),
        ("Base destino mais afetada · 目的网点", f"{n['base_destino'][0]} ({n['base_destino'][1]})"),
        ("Segmento mais afetado · 二段码", f"{n['segmento'][0]} ({n['segmento'][1]})"),
        ("ID de viagem mais afetado · 车次号", f"{n['id_viagem'][0]} ({n['id_viagem'][1]})"),
        ("Lote (saca) mais afetado · 包号", f"{n['lote'][0]} ({n['lote'][1]})"),
        ("Base destino fora da rotinha · 未匹配网点", n["fora_rotinha"]), ("Rota fora da tabela · 未匹配线路", n["rota_fora"]),
    ]
    for rot, val in pares:
        fmt = "0.0%" if isinstance(val, float) else None
        escrever(rot, _cel(ws, val, Font(name=_FONTE, bold=True), formato=fmt))
    escrever()

    errados = pacotes[pacotes["situacao"] == config.ERRADO] if not pacotes.empty else pacotes
    blocos = [
        ("Errados por turno · 各班次错发", contar(errados["turno"], ordem=list(config.TURNOS)), "rosca"),
        ("Base de escaneamento · 扫描网点", contar(errados["base_escaneamento"], topo=10), "barra"),
        ("Bases de destino mais afetadas · 目的网点", contar(errados["base_destino"], topo=15), "barra"),
        ("Segmentos mais afetados (Código de dois segmentos) · 二段码", contar(errados["segmento"], topo=15), "barra"),
        ("Destinos (Nome de mecanismo final) · 最终机构", contar(errados["destino"], topo=15), "barra"),
        ("Logins (Digitalizado por) · 扫描人", contar(errados["digitalizado_por"], topo=15), "barra"),
        ("ID de viagem mais afetado · 车次号", contar(errados["id_viagem"], topo=15), "barra"),
        ("Número do lote (sacas com mais erro) · 包号", contar(errados["lote"], topo=15), "barra"),
        ("Situação de todos os pacotes · 判定", contar(pacotes["situacao"], ordem=list(config.CORES_SITUACAO)), "situacao"),
    ]
    for titulo, df, forma in blocos:
        escrever(_cel(ws, titulo, _SUB))
        ws.append([_cel(ws, "Valor", _BRANCO, _CAB_PT, _CENTRO), _cel(ws, "Quantidade", _BRANCO, _CAB_PT, _CENTRO),
                   _cel(ws, "% do total", _BRANCO, _CAB_PT, _CENTRO)])
        ws.append([_cel(ws, "值", _BRANCO_ZH, _CAB_ZH, _CENTRO), _cel(ws, "数量", _BRANCO_ZH, _CAB_ZH, _CENTRO),
                   _cel(ws, "占比", _BRANCO_ZH, _CAB_ZH, _CENTRO)])
        linha += 2
        topo_bloco = linha - 2
        primeira = linha + 1
        for r in df.itertuples(index=False):
            escrever(r.valor, int(r.quantidade), _cel(ws, float(r.percentual) / 100, formato="0.0%"))
        if df.empty:
            escrever("Sem pacotes errados", 0, 0)
        ultima = linha
        ancora = f"E{topo_bloco}"
        if forma == "rosca":
            _grafico_rosca(ws, titulo, primeira, ultima, ancora, [config.CORES_TURNO.get(v, config.COR_NEUTRA) for v in df["valor"]])
        elif forma == "situacao":
            _grafico_barras(ws, titulo, 1, 2, primeira, ultima, ancora,
                            [config.CORES_SITUACAO.get(v, config.COR_NEUTRA) for v in df["valor"]], horizontal=False)
        else:
            _grafico_barras(ws, titulo, 1, 2, primeira, ultima, ancora)
        # Espaço para o gráfico (cerca de 15 linhas) antes do próximo bloco.
        for _ in range(max(2, 16 - (ultima - topo_bloco))):
            escrever()

    _aba_dados(wb, "Errados 错发明细", errados, config.COLUNAS_PACOTES)
    _aba_dados(wb, "Viagens 车次", viagens, config.COLUNAS_VIAGENS)
    _aba_dados(wb, "Todos os pacotes 全部", pacotes, config.COLUNAS_PACOTES)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
