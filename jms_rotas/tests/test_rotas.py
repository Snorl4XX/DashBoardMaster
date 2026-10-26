import io
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rotaslib import analise, config, tabelas  # noqa: E402
from rotaslib.cliente import ClienteJMS, ErroJMS, linha_de_carregamento  # noqa: E402
from rotaslib.demo import JMSFicticio, rotas_demo, rotinha_demo  # noqa: E402
from rotaslib.monitor import Ajustes, Monitor  # noqa: E402
from rotaslib.planilha import gerar_planilha  # noqa: E402


class Resposta:
    def __init__(self, status, corpo):
        self.status_code = status
        self.text = corpo if isinstance(corpo, str) else json.dumps(corpo, ensure_ascii=False)
        self.encoding = "utf-8"


class SessaoFalsa:
    def __init__(self, responder):
        self.responder = responder
        self.pedidos = []

    def get(self, url, params=None, headers=None, timeout=None, allow_redirects=True):
        self.pedidos.append(("GET", url, dict(params or {}), dict(headers)))
        return Resposta(*self.responder("GET", url, params or {}, None, headers))

    def post(self, url, data=None, headers=None, timeout=None, allow_redirects=True):
        corpo = json.loads(data)
        self.pedidos.append(("POST", url, corpo, dict(headers)))
        return Resposta(*self.responder("POST", url, {}, corpo, headers))


def _ok(data):
    return 200, {"code": 1, "msg": "请求成功", "data": data, "succ": True, "fail": False}


# Linhas do bloco "carga e descarga" da secundária (como no PDF, parte 06).
CARGAS_SEC = [
    {"jobCode": "SRTR22606546911", "loadingTypeName": "2", "scanNetworkName": "F S-SAP 02-SP", "scanNetworkCode": "311614",
     "scanWaybillNum": 12},
    {"jobCode": "SRTR22606546911", "loadingTypeName": "1", "scanNetworkName": "SP GRU", "scanNetworkCode": "30001",
     "scanWaybillNum": 12},
]


def _pacote(i, destsite, proxima="F S-SAP 02-SP"):
    return {"billCode": f"88800275{i:07d}", "networkName": "SP GRU", "networkCode": "30001", "nextStation": proxima,
            "packageCode": "BR10506244314", "scanTime": "2026-10-26 04:40:28", "scanType": "1", "scanUser": "PESSOA TESTE",
            "total": "12", "destsite": destsite, "destsiteCity": "São Paulo"}


def test_linha_de_carregamento_e_a_da_base():
    assert linha_de_carregamento(CARGAS_SEC, "30001", "SP GRU")["scanNetworkCode"] == "30001"
    assert linha_de_carregamento([CARGAS_SEC[0]], "30001", "SP GRU") is None
    so_nome = [{**CARGAS_SEC[1], "scanNetworkCode": ""}]
    assert linha_de_carregamento(so_nome, "30001", "SP GRU")["scanNetworkName"] == "SP GRU"


def test_viagens_troncal_e_secundaria_com_os_filtros_do_jms():
    def responder(metodo, url, params, corpo, headers):
        if url == config.URLS[config.TRONCAL]["viagens"]:
            return _ok({"records": [{"shipmentNo": "DKGX1", "startCode": "30001", "shipmentState": 4},
                                    {"shipmentNo": "DKGX2", "startCode": "62101", "shipmentState": 4}], "total": 2, "pages": 1})
        return _ok({"records": [{"shipmentNo": "SRTR1", "startCode": "30001", "startName": "SP GRU", "shipmentState": 3},
                                {"shipmentNo": "SRTR2", "startCode": "311036", "startName": "DC SBN-SP", "shipmentState": 4}],
                    "total": 2, "pages": 1})
    s = SessaoFalsa(responder)
    c = ClienteJMS("tok", sessao=s, espera_base=0)
    assert [v["shipmentNo"] for v in c.viagens(config.TRONCAL, "2026-10-26", "30001")] == ["DKGX1"]
    assert [v["shipmentNo"] for v in c.viagens(config.SECUNDARIA, "2026-10-26", "30001")] == ["SRTR1"]
    metodo, url, params, cab = s.pedidos[0]
    assert metodo == "GET" and params["startCode"] == "30001" and params["searchType"] == "manage"
    assert params["startDateTime"] == "2026-10-26 00:00:00" and params["endDateTime"] == "2026-10-26 23:59:59"
    assert cab["AuthToken"] == "tok"
    metodo, url, corpo, _ = s.pedidos[1]
    assert metodo == "POST" and corpo["startDepartureTime"] == "2026-10-26 00:00:00" and corpo["countryId"] == "1"


def test_pacotes_paginados_e_pagina_cortada_pelo_jms():
    todos = [_pacote(i, "F S-SAP 02-SP") for i in range(45)]

    def responder(metodo, url, params, corpo, headers):
        assert url == config.URLS[config.TRONCAL]["pacotes"]
        assert params["scanNetworkCode"] == "30001" and params["countNum"] == 45
        # O JMS corta o tamanho em 20, qualquer que seja o pedido.
        ini = (params["current"] - 1) * 20
        return _ok({"records": todos[ini:ini + 20], "total": 45, "size": params["size"], "current": params["current"]})
    s = SessaoFalsa(responder)
    lista = ClienteJMS("tok", sessao=s, espera_base=0).pacotes(config.TRONCAL, "DBGX1", "30001", 45)
    assert len(lista) == 45 and len({p["billCode"] for p in lista}) == 45
    assert len(s.pedidos) == 3


def test_token_vencido():
    s = SessaoFalsa(lambda *a: (200, {"code": 135010037, "msg": "token失效，请重新登录", "fail": True}))
    with pytest.raises(ErroJMS) as e:
        ClienteJMS("tok", sessao=s, espera_base=0).cargas(config.SECUNDARIA, "SRTR1")
    assert e.value.tipo == "auth" and len(s.pedidos) == 1


def test_regra_com_tabela_de_rotas():
    ref = analise.Referencias(
        rotinha={tabelas.chave("OUR -PE"): ("OUR -PE", "551-01", "DC SNS-PE"), tabelas.chave("SDR-SP"): ("SDR-SP", "375-02", "DC SBN-SP"),
                 tabelas.chave("F S-SAP 02-SP"): ("F S-SAP 02-SP", "120-02", "F S-SAP 02-SP")},
        rotas={tabelas.chave("BA FEC"): {tabelas.chave("BA FEC"), tabelas.chave("DC SNS-PE")}})
    assert analise.classificar("OUR -PE", "BA FEC", ref)[2] == config.CORRETO
    assert analise.classificar("SDR-SP", "BA FEC", ref)[2] == config.ERRADO          # pacote de SP no caminhão da BA
    assert analise.classificar("XYZ-SP", "BA FEC", ref)[2] == config.SEM_ROTINHA      # #N/A no PROCV
    assert analise.classificar("SDR-SP", "SP BRE", ref)[2] == config.ROTA_FORA       # rota que não está na tabela
    assert analise.classificar("F S-SAP 02-SP", "F S-SAP 02-SP", ref)[2] == config.CORRETO
    seg, dest, _, _ = analise.classificar("OUR -PE", "BA FEC", ref)
    assert (seg, dest) == ("551-01", "DC SNS-PE")


def test_regra_sem_tabela_de_rotas_e_segmento_na_tabela():
    ref = analise.Referencias(rotinha={tabelas.chave("S-SRAF-SP"): ("S-SRAF-SP", "121-07", "S-SRAF-SP")})
    assert analise.classificar("S-SRAF-SP", "F S-SAP 02-SP", ref)[2] == config.ERRADO
    assert analise.classificar("S-SRAF-SP", "S-SRAF-SP", ref)[2] == config.CORRETO
    ref.rotas = {tabelas.chave("F S-SAP 02-SP"): {"121-07"}}  # tabela de rotas com segmentos
    assert analise.classificar("S-SRAF-SP", "F S-SAP 02-SP", ref)[2] == config.CORRETO


def _xlsx(linhas, aba="Sheet0"):
    wb = Workbook()
    wb.active.title = aba
    for l in linhas:
        wb.active.append(l)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_rotinha_com_cabecalho_e_sem_cabecalho():
    r = tabelas.ler_rotinha("rotinhas.xlsx", _xlsx([["Nome da base", "Código de dois segmentos", "Nome de mecanismo final", "Estado"],
                                                    ["CANA -BA", "512-00", "BA FEC", "BA"], ["CANA -BA", "999", "X", "BA"]]))
    assert r[tabelas.chave("CANA-BA")] == ("CANA -BA", "512-00", "BA FEC")  # como o PROCV: a primeira linha
    r = tabelas.ler_rotinha("r.csv", "F IPR-GO;601-01;GO GYN\nPER -BA;507-00;BA FEC\n".encode())
    assert r[tabelas.chave("PER -BA")][1:] == ("507-00", "BA FEC")


def test_tabela_de_rotas_nos_dois_formatos():
    linhas = _xlsx([["Rota", "Destinos"], ["BA FEC", "BA FEC, BA SBA"], ["", "DC SNS-PE"], ["SP BRE", "SP BRE"]])
    rotas, sug = tabelas.ler_rotas("rotas.xlsx", linhas)
    assert sug["modo"] == tabelas.MODO_LINHAS
    assert rotas[tabelas.chave("BA FEC")] == {"BA FEC", "BA SBA", "DC SNS-PE"}  # célula mesclada: continua na rota de cima
    colunas = _xlsx([["BA FEC", "SP BRE"], ["BA SBA", "DC SBN-SP"], ["DC SNS-PE", ""]])
    rotas, sug = tabelas.ler_rotas("rotas.xlsx", colunas)
    assert sug["modo"] == tabelas.MODO_COLUNAS
    assert rotas[tabelas.chave("SP BRE")] == {"DC SBN-SP"}


def test_monitor_le_so_concluido_e_em_transito_e_nao_repete(tmp_path):
    m = Monitor(tmp_path)
    cli = JMSFicticio()
    m.cliente, m.ajustes = cli, Ajustes(dia="2026-10-26")
    m.rodada(config.TRONCAL)
    f = m.foto(config.TRONCAL)
    assert f["viagens"] and all(int(v["shipmentState"]) in (3, 4) for v in f["viagens"].values())
    assert not any(i.endswith("99998") or i.endswith("99999") for i in f["viagens"])  # aguardando / outra base
    assert set(f["pacotes"]) == set(f["viagens"])
    antes = cli.consultas
    m.rodada(config.TRONCAL)
    em_transito = sum(1 for v in f["viagens"].values() if int(v["shipmentState"]) == 3)
    assert cli.consultas - antes == 1 + em_transito  # lista + só as em trânsito; concluídas não são consultadas de novo
    # Cache em disco: outro monitor abre já com os pacotes.
    m2 = Monitor(tmp_path)
    m2.configurar(None, Ajustes(dia="2026-10-26"))
    assert set(m2.foto(config.TRONCAL)["pacotes"]) == set(f["pacotes"])
    m2.parar()


def test_monitor_para_com_token_vencido(tmp_path):
    m = Monitor(tmp_path)
    m.cliente, m.ajustes = JMSFicticio(token_recusado=True), Ajustes(dia="2026-10-26")
    m.rodada(config.SECUNDARIA)
    assert "AuthToken" in m.foto(config.SECUNDARIA)["erro"] and m.pausado_por


def test_cartoes_e_planilha(tmp_path):
    m = Monitor(tmp_path)
    m.cliente, m.ajustes = JMSFicticio(), Ajustes(dia="2026-10-26")
    m.rodada(config.SECUNDARIA)
    f = m.foto(config.SECUNDARIA)
    df = analise.tabela_pacotes(f, analise.Referencias(rotinha_demo(), rotas_demo()))
    vg = analise.tabela_viagens(f, df)
    n = analise.cartoes(df, vg)
    assert n["errados"] == int((df["situacao"] == config.ERRADO).sum()) > 0
    assert round(sum(n["turnos"][t]["participacao"] for t in config.TURNOS)) == 100
    assert n["viagens_concluidas"] + n["viagens_transito"] == len(vg)
    assert set(vg["errados"]) and vg["errados"].sum() == n["errados"]
    wb = load_workbook(io.BytesIO(gerar_planilha(config.SECUNDARIA, df, vg, n, {"dia": "2026-10-26"})))
    assert wb.sheetnames == ["Resumo 汇总", "Errados 错发明细", "Viagens 车次", "Todos os pacotes 全部"]
    ws = wb["Errados 错发明细"]
    assert ws["A1"].value == "ID de Viagem" and ws["A2"].value == "车次号" and ws["A1"].fill.fgColor.rgb.endswith("E60012")
    assert ws.max_row == n["errados"] + 2
    assert [c.value for c in ws[1]][12:15] == ["Turno", "SEGMENTOS (Código de dois segmentos)", "DESTINOS (Nome de mecanismo final)"]
    assert len(wb["Resumo 汇总"]._charts) == 9
