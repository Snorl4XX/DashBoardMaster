"""Testes do Expedição Incorreta (painel Streamlit). Rodar:  py -m pytest tests"""

from __future__ import annotations

import io
import json
import os
import sys
import threading
from datetime import date, datetime, time as dt_time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

PASTA = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PASTA))

from expedicao_dash import analise, demo, planilha, relatorio  # noqa: E402
from expedicao_dash.cache import Cache  # noqa: E402
from expedicao_dash.jms import Acesso, Cliente, Consulta, Demo, consultar_lote, ler_resposta  # noqa: E402
from expedicao_dash.rotas import ler_rotas, tabela_padrao  # noqa: E402
from expedicao_dash.texto import (data_hora, dividir_terminal, lista_primeiros, lista_segundos, norm_remessa, norm_segundo,  # noqa: E402
                                  so_hora, turno_da_hora)

# ------------------------------------------------------------------ a planilha de rotas no formato das tabelas vermelhas
REGIONAIS = [
    ("6", "01", "DC S-VLANA-SP", "VLANA", None, "预留Reservado"), (None, "02", "F S-JBO-SP", "GRU", "404-11", "预留Reservado"),
    (None, "03", "F S-SAP 02-SP", "GRU", "404-12", "预留Reservado"), (None, "04", "F S-MGUE 02-SP", "GRU", "404-08", "有立柱"),
    (None, "05", "F S-PQBOTU-SP", "GRU", "404-09", None), (None, "06", "F VLMT-SP", "GRU", "404-10", None),
    (None, "07", "S-MGUE-SP", "GRU", "404-01", None), (None, "08", "DC SSZ-SP", "GRU", "378-00、01、02、03、05、06", "有立柱"),
    (None, "09", None, None, None, None), (None, "10", "S-MGUE 02-SP", "GRU", "404-13", "预留Reservado"),
    ("5", "11", "DC S-MOO 02-SP", "SP", "（Antigo）364-00、435-00、350-00、365-00、381-12、430—01、05、07、08、09", "重新调整Reajustado"),
    (None, "12", None, "MOO", "（Novo）430—00、01、02、03、04、05、06、07、08、09", None), (None, "13", "S-PENHA-SP", "GRU", "404-02", None),
    (None, "14", "F S-PENHA 02-SP", "GRU", "404-07", "有立柱"), (None, "15", "S-SRAF-SP", "GRU", "404-00", None),
    (None, "16", None, None, None, None), (None, "17", "DC NOD-SP", "NOD", None, "预留Reservado"), (None, "18", None, None, None, None),
    (None, "19", "S-SAPOP-SP", "GRU", "404-03", None), (None, "20", "F S-MGUE-SP", "GRU", "404-04", "有立柱"), (None, "21", "-", "-", "-", None),
    (None, "22", "F S-PENHA-SP", "GRU", "404-05", None), (None, "23", "-", "-", "-", None), (None, "24", "F S-PENHA 03-SP", "GRU", "404-06", None),
    (None, "25", "-", "-", "-", None),
]
LINHA = [
    ("7", "05", "保持原码头", "GO GYN", "GO", None, None), (None, "06", None, "SC BNU", "SC、SC1", None, None),
    (None, "07", None, "RS NSR", "RS、RS1", None, None), (None, "08", None, "SBN", "SBN", "375—00、01、02、03、04、05、06、07、08、09、10", None),
    (None, "09", None, "CE FOR", "CE", None, None), (None, "10", None, "RJ SJM", "RJ、RJ1、VDC、CDG", None, "新增一段码CDG"),
    ("4", "15", "保持原码头", "SE AJU", "SE、AL", None, None), (None, "16", None, "PR SJS", "PR、PR1、CAC", None, None),
    (None, "17", None, "PE JGS", "PE、DC—BYE", None, None), (None, "18", None, "BA FEC", "BA、SBA、SNS", None, None),
    (None, "19", None, "MG CGE", "MG、MG1、CHV", None, None), (None, "20", None, "PA ANA", "PA", None, "人工拉包"),
    ("2", "21", "保持原码头", "SP BRE", "AC、AM、BAU、BVB、JDF、LDB、MIA、MRB、RO、SJP、SOD、SP、SP1、STM、TO、VCP、XAP、BJE、MA、PPB、NAT", None, "Excluir CDG"),
    (None, "22", None, "SP BRE", "AC、AM、BAU、BVB、JDF、LDB、MIA、MRB、RO、SJP、SOD、SP、SP1、STM、TO、VCP、XAP、BJE、MA、PPB、NAT", None, "Excluir CDG"),
    (None, "23", "14", "DC TBT-SP", "SP（Antigo）", "402-01、03、05\n421-00", "Transferido da doca nº 14 para a doca nº 23."),
    (None, None, None, None, "GRU（Novo）", "402—00、01、03、05、06、07、08\n402—02、04", None),
    ("3", "44", "保持原码头", "DC GRU-SP", "SP（Antigo）", "395—05、06、07\n381—04", "378-04 — Data de ativação a definir."),
    (None, None, None, None, "GRU（Novo）", "381—00、01、02、03、04、05、11、13、14、15、18、19、20、21\n378—04", None),
    (None, "43", None, "MS CGR", "MS", None, None), (None, "42", None, "MT CGB", "MT、JPN", None, None), (None, "41", None, "THE", "PI、IMP", None, None),
    (None, "40", None, "ES SRR", "ES", None, None), (None, "39", None, "SP RBP", "NE、NE1", None, None), (None, "38", None, "DF BSB", "DF", None, None),
]


def planilha_rotas_24_10() -> bytes:
    """Mesmo formato da ajuste_rotas_docas_24_outubro.xlsx: título mesclado, cabeçalho na linha 2, rotas mescladas em 2 docas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rotas Regionais"
    ws.append(["10月24日支线装车码头进行调整（更新10.24）\nAjuste das Docas de Carregamento das Rotas Regionais em 24 de Outubro"])
    ws.append(["DWS摆轮线路\nLinha de desvio do DWS", "新码头号码\nNovo número da doca", "调整后的装车线路\nRota de Carregamento Após o Ajuste",
               "一段码\n1º Nível", "二段码\n2º Nível", "备注\nObservações"])
    for linha in REGIONAIS:
        ws.append(list(linha))
    for faixa in ("A1:F1", "C10:C11", "D10:D11", "E10:E11", "C13:C14", "C17:C18", "C19:C20"):
        ws.merge_cells(faixa)
    ws2 = wb.create_sheet("Rotas de Linha")
    ws2.append(["10月24日干线装车码头进行调整（更新10.24）"])
    ws2.append(["DWS摆轮线路\nLinha de desvio do DWS", "新码头号码\nNovo número da doca", "原码头号码\nNúmero original da doca",
                "调整后的装车线路\nRota de Carregamento Após o Ajuste", "段码\n1º Nível", "二段码\n2º Nível", "备注\nObservações"])
    for linha in LINHA:
        ws2.append(list(linha))
    for faixa in ("A1:G1", "D17:D18", "D19:D20", "C17:C18"):
        ws2.merge_cells(faixa)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _chave(tab):
    return sorted((r.rota, r.primeiros, tuple(sorted(r.segundos)), r.tipo) for r in tab.regras)


# ------------------------------------------------------------------ texto
def test_textos_e_codigos():
    assert norm_remessa(888002560356352.0) == "888002560356352"
    assert norm_remessa("888002560356352.0") == "888002560356352"
    assert norm_remessa("8.88002560356352E+14") == "888002560356352"
    assert norm_remessa(" 888002223648620 ") == "888002223648620"
    assert lista_segundos("（Antigo）364-00、435-00、430—01、05") == ["364-00", "435-00", "430-01", "430-05"]
    assert lista_segundos("402-01、03、05\n421-00") == ["402-01", "402-03", "402-05", "421-00"]
    assert lista_segundos("404 - 05") == ["404-05"]
    assert lista_primeiros("PE、DC—BYE") == ["PE", "DC-BYE"] and lista_primeiros("SP（Antigo）") == ["SP"] and lista_primeiros("-") == []
    assert dividir_terminal("SP,435-00,080") == ("SP", "435-00", "080")
    assert dividir_terminal("MRB，742-00，001") == ("MRB", "742-00", "001")
    assert dividir_terminal("GO") == ("GO", "", "")  # só o 1º segmento também serve (troncal)
    assert norm_segundo("40405") == "404-05"
    assert data_hora("2026-10-27 14:05:00") == datetime(2026, 10, 27, 14, 5)
    assert data_hora("27/10/2026 07:10") == datetime(2026, 10, 27, 7, 10)
    assert data_hora(46322.5) == datetime(2026, 10, 27, 12, 0)
    assert so_hora(0.5) == dt_time(12, 0) and so_hora("22:15") == dt_time(22, 15)
    assert [turno_da_hora(dt_time(h)) for h in (5, 6, 13, 14, 21, 22)] == ["T3", "T1", "T1", "T2", "T2", "T3"]


# ------------------------------------------------------------------ planilha de rotas
def test_planilha_de_rotas_no_formato_das_docas():
    t = ler_rotas(planilha_rotas_24_10(), "ajuste_rotas_docas_24_outubro.xlsx")
    # O app antigo lia a aba "Rotas Regionais" com as colunas da troncal (e criava rotas chamadas GRU, SP, MOO).
    assert {"GRU", "SP", "MOO"}.isdisjoint(t.rotas)
    assert _chave(t) == _chave(tabela_padrao())
    assert len(t.rotas) == 39 and not t.avisos
    ssz = next(r for r in t.regras if r.rota == "DC SSZ-SP")
    assert ssz.docas == "08, 09" and ssz.segundos == ("378-00", "378-01", "378-02", "378-03", "378-05", "378-06")
    real = os.environ.get("ROTAS_REAIS")  # a planilha de verdade (fora do git), se informada
    if real and Path(real).exists():
        assert _chave(ler_rotas(Path(real).read_bytes(), Path(real).name)) == _chave(tabela_padrao())


@pytest.mark.parametrize("p,s,rota", [
    ("GRU", "404-05", "F S-PENHA-SP"),  # exemplo do documento (secundária: par)
    ("SP", "435-00", "DC S-MOO 02-SP"),  # SP antigo da DC S-MOO
    ("SP", "999-00", "SP BRE"),  # SP sem segmento especial: troncal pelo 1º
    ("AC", "", "SP BRE"),  # "AC pertence a BRE"
    ("SP", "402-01", "DC TBT-SP"), ("GRU", "402-00", "DC TBT-SP"),  # "GRU 402-00 é para a rota DC TBT-SP"
    ("GRU", "381-18", "DC GRU-SP"), ("SBN", "375-04", "SBN"), ("MOO", "430-03", "DC S-MOO 02-SP"),
    ("VLANA", "", "DC S-VLANA-SP"), ("NOD", "001-00", "DC NOD-SP"), ("BYE", "100-01", "PE JGS"), ("CDG", "", "RJ SJM"),
])
def test_rota_certa(p, s, rota):
    e = tabela_padrao().esperada(p, s)
    assert (e.rota, e.situacao) == (rota, "ok")


def test_rota_sem_resposta_e_nomes():
    t = tabela_padrao()
    assert t.esperada("GRU", "999-99").situacao == "segmento_fora"
    assert t.esperada("XYZ", "").situacao == "codigo_fora"
    assert t.esperada("", "").situacao == "sem_codigo"
    assert t.resolver("PI THE") == "THE" and t.resolver("DC SBN-SP") == "SBN"  # nomes do JMS × nomes da tabela
    assert t.resolver("dc s-moo02-sp") == "DC S-MOO 02-SP"
    assert t.resolver("F S-PENHA 04-SP") == "" and t.resolver("SP BREE") == ""  # outra rota / nome errado: não adivinha


def test_planilha_de_rotas_simples():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tipo", "Rota (Próxima parada)", "1º Nível (códigos)", "2º Nível (segmentos, opcional)", "Observação"])
    ws.append(["Troncal", "SP BRE", "SP, AC", "", ""])
    ws.append(["Secundária", "F S-PENHA-SP", "GRU", "404-05", ""])
    b = io.BytesIO()
    wb.save(b)
    t = ler_rotas(b.getvalue(), "rotas.xlsx")
    assert t.esperada("GRU", "404-05").rota == "F S-PENHA-SP" and t.esperada("AC").rota == "SP BRE"
    with pytest.raises(ValueError):
        wb2 = openpyxl.Workbook()
        b2 = io.BytesIO()
        wb2.save(b2)
        ler_rotas(b2.getvalue(), "vazia.xlsx")


# ------------------------------------------------------------------ planilha de remessas
def test_ler_planilha_de_remessas():
    csv_ = "Relatório de carga\n\nNúmero de pedido JMS;Próxima parada;Data;Hora\n888002560356352;SP BRE;27/10/2026;14:05\n" \
           "888002560356352;SP BRE;27/10/2026;14:05\n888002223648620;DC S-MOO 02-SP;27/10/2026;06:30\nxx;SP BRE;;\n"
    abas = planilha.abrir(csv_.encode("cp1252"), "lista.csv")
    df = abas["CSV"]
    cols = planilha.sugerir_colunas(df)
    assert (cols["remessa"], cols["rota"], cols["data"], cols["hora"]) == ("Número de pedido JMS", "Próxima parada", "Data", "Hora")
    l = planilha.montar(df, cols)
    assert l.tabela["remessa"].tolist() == ["888002560356352", "888002223648620"] and l.repetidas == 1 and l.invalidas == 1
    assert l.tabela["turno"].tolist() == ["T2", "T1"] and l.tabela["data"].tolist() == [date(2026, 10, 27)] * 2
    # .xlsx com título em cima, remessa como número e data/hora juntas.
    b = io.BytesIO()
    with pd.ExcelWriter(b, engine="openpyxl") as w:
        pd.DataFrame([["Lista de carga"], [None]]).to_excel(w, index=False, header=False, sheet_name="rj")
        pd.DataFrame({"Remessa": [888002560356352, 888002223648620], "Tempo de digitalização": ["2026-10-27 23:10:00", "2026-10-27 15:00:00"],
                      "Rota": ["SP BRE", "PI THE"]}).to_excel(w, index=False, startrow=2, sheet_name="rj")
    df = planilha.abrir(b.getvalue(), "x.xlsx")["rj"]
    cols = planilha.sugerir_colunas(df)
    l = planilha.montar(df, cols, viagem_fixa="T123")
    assert l.tabela["remessa"].tolist() == ["888002560356352", "888002223648620"]
    assert l.tabela["turno"].tolist() == ["T3", "T2"] and set(l.tabela["viagem"]) == {"T123"}
    assert planilha.montar(df, cols, rota_fixa="SP BRE").tabela["rota"].tolist() == ["SP BRE", "SP BRE"]


# ------------------------------------------------------------------ JMS
RESPOSTA_PDF = {"code": 1, "msg": "1:Solicitação concluída", "succ": True, "fail": False, "data": {"waybillNo": "888002560356352", "details": {
    "terminalDispatchCode": "SP,435-00,080", "realPickNetworkName": "PA MELI-VCP-SP", "senderCityName": "Campinas", "initDistributeName": "SP BRE",
    "dispatchNetworkName": "F S-VLMR-SP", "receiverCityName": "São Paulo", "destinationDistributeName": "SP GRU", "receiverName": "******",
    "receiverMobilePhone": "*******6869", "receiverTaxNumber": "07809796801"}}}


def test_ler_resposta_do_jms():
    c = ler_resposta("888002560356352", 200, json.dumps(RESPOSTA_PDF))
    assert c.ok and (c.primeiro, c.segundo, c.terceiro) == ("SP", "435-00", "080")
    assert c.detalhes == {"PDD de saída": "PA MELI-VCP-SP", "Município de Origem": "Campinas", "SC Origem": "SP BRE",
                          "Base Destino": "F S-VLMR-SP", "Município de Destino": "São Paulo", "SC Destino": "SP GRU"}
    assert "07809796801" not in c.para_json() and "6869" not in c.para_json()  # dados pessoais não são guardados
    # Token vencido com HTTP 200 (o app antigo não percebia e seguia o lote inteiro).
    assert ler_resposta("1", 200, json.dumps({"code": 135010035, "msg": "登录失效，请重新登录"})).tipo_erro == "auth"
    assert ler_resposta("1", 401, "").tipo_erro == "auth"
    assert ler_resposta("1", 200, "<html><title>Login JMS</title></html>").tipo_erro == "auth"
    assert ler_resposta("1", 200, "<html><title>Access denied</title></html>").tipo_erro == "bloqueio"
    assert ler_resposta("1", 200, json.dumps({"code": 0, "msg": "请求过于频繁"})).tipo_erro == "limite"
    assert ler_resposta("1", 429, "").tipo_erro == "limite" and ler_resposta("1", 502, "").tipo_erro == "rede"
    assert ler_resposta("1", 200, json.dumps({"code": 1, "data": None})).tipo_erro == "nao_achou"
    sem = ler_resposta("1", 200, json.dumps({"code": 1, "data": {"details": {"terminalDispatchCode": None}}}))
    assert sem.ok and sem.tipo_erro == "sem_codigo"
    so1 = ler_resposta("1", 200, json.dumps({"code": 1, "data": {"details": {"terminalDispatchCode": "GO"}}}))
    assert so1.ok and so1.primeiro == "GO" and not so1.tipo_erro


class _JmsFalso:
    """Servidor local que responde como o getOrderDetail e guarda os cabeçalhos recebidos."""

    def __init__(self, token="tok", vence_depois=None):
        self.token, self.vence_depois, self.pedidos, self.cabecalhos = token, vence_depois, 0, []
        dono = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_POST(self):
                corpo = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                dono.pedidos += 1
                dono.cabecalhos.append(dict(self.headers))
                if self.headers.get("authtoken") != dono.token or (dono.vence_depois and dono.pedidos > dono.vence_depois):
                    r = {"code": 135010035, "msg": "token失效"}
                else:
                    w = corpo["waybillNo"]
                    r = {"code": 1, "data": {"waybillNo": w, "details": {"terminalDispatchCode": f"GRU,404-0{int(w[-1]) % 7},010",
                                                                          "dispatchNetworkName": "F S-PENHA-SP"}}}
                dados = json.dumps(r).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(dados)))
                self.end_headers()
                self.wfile.write(dados)

        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.srv.server_address[1]}/operatingplatform/order/getOrderDetail"


def test_cliente_e_lote_no_jms_falso():
    jms = _JmsFalso()
    cli = Cliente(Acesso(token="tok", url=jms.url))
    ws = [f"88800256035{i:04d}" for i in range(30)]
    prontas, a = consultar_lote(ws, cli.consultar, paralelo=5)
    assert a.feitas == 30 and not a.parou_auth and {c.remessa for c in prontas} == set(ws)
    assert all(c.ok and c.primeiro == "GRU" for c in prontas)
    h = jms.cabecalhos[0]
    assert h["Routename"] == "trackingExpress" and h["Lang"] == "PT" and h["Routernamelist"].startswith("%E6%93%8D")
    # Token vence no meio: para na hora (não consulta as 200 restantes) e devolve o que já veio.
    jms2 = _JmsFalso(vence_depois=12)
    prontas, a = consultar_lote([f"88800256036{i:04d}" for i in range(200)], Cliente(Acesso(token="tok", url=jms2.url)).consultar, paralelo=4)
    assert a.parou_auth and 8 <= len(prontas) <= 12 and jms2.pedidos < 40 and "AuthToken" in a.mensagem


def test_lote_espera_quando_o_jms_pede():
    vezes = {}

    def consultar(w):
        vezes[w] = vezes.get(w, 0) + 1
        if vezes[w] == 1 and w.endswith("1"):
            return Consulta(w, erro="muitas consultas", tipo_erro="limite")
        return Consulta(w, ok=True, primeiro="SP", segundo="435-00")
    esperas = []
    prontas, a = consultar_lote([f"88800256000000{i}" for i in range(10)], consultar, paralelo=4, esperar=esperas.append)
    assert len(prontas) == 10 and all(c.ok for c in prontas) and esperas == [3] and a.paralelo == 2


def test_cache(tmp_path):
    c = Cache(tmp_path / "c.sqlite3")
    ws = [f"8880{i:011d}" for i in range(1200)]  # mais que o limite de "?" do SQLite antigo
    c.guardar([Consulta(w, True, "GRU", "404-05", "010", {"SC Destino": "SP GRU"}) for w in ws] +
              [Consulta("x1", True, erro="sem", tipo_erro="sem_codigo"), Consulta("x2", erro="rede", tipo_erro="rede")])
    lidas = c.ler(ws + ["x1", "x2"])
    assert len(lidas) == 1201 and lidas[ws[5]].detalhes == {"SC Destino": "SP GRU"} and "x2" not in lidas
    assert lidas["x1"].tipo_erro == "sem_codigo"
    assert c.ler(["x1"], horas_sem_codigo=0) == {}
    c.apagar()
    assert c.quantas() == 0


# ------------------------------------------------------------------ análise e Excel
def _resultado():
    linhas, gabarito = demo.gerar()
    df = planilha.abrir(planilha.exemplo_planilha(linhas), "exemplo.xlsx")["Info de carga e descarga"]
    tab = planilha.montar(df, planilha.sugerir_colunas(df)).tabela
    prontas, _ = consultar_lote(tab["remessa"].tolist(), Demo(gabarito).consultar, paralelo=8)
    res = analise.classificar(analise.juntar_jms(tab, {c.remessa: c for c in prontas}), tabela_padrao())
    return res, analise.resumir(res)


def test_analise_caso_do_documento():
    tab = pd.DataFrame([{"remessa": "00000000000", "rota": "SP BRE", "viagem": "", "quando": None, "data": None, "hora": "14:00",
                         "turno": "T2", "primeiro": "", "segundo": ""},
                        {"remessa": "00000000001", "rota": "DC S-MOO 02-SP", "viagem": "", "quando": None, "data": None, "hora": "",
                         "turno": "Sem horário", "primeiro": "", "segundo": ""}])
    cons = {"00000000000": Consulta("00000000000", True, "GRU", "404-05"), "00000000001": Consulta("00000000001", True, "SP", "435-00")}
    res = analise.classificar(analise.juntar_jms(tab, cons), tabela_padrao())
    assert res["situacao"].tolist() == ["Incorreta", "Correta"]
    assert res.loc[0, "rota_certa"] == "F S-PENHA-SP" and "F S-PENHA-SP" in res.loc[0, "motivo"]


def test_resumo_e_excel():
    res, r = _resultado()
    assert r.total == len(res) and r.analisadas + r.pendentes == r.total and r.incorretas > 20
    assert r.taxa == pytest.approx(r.incorretas / r.analisadas)
    assert r.por_turno["Incorretas"].sum() == r.incorretas and r.por_rota_carregada["Incorretas"].sum() == r.incorretas
    assert r.por_segmento["Incorretas"].sum() == r.incorretas and r.por_dia["Incorretas"].sum() == r.incorretas
    assert abs(r.por_rota_certa["% dos erros"].sum() - 1) < 1e-9
    assert set(r.pendencias["Motivo"]) >= {analise.P_JMS_SEM, analise.P_NAO_ACHOU}
    x = relatorio.gerar_excel(res, r, tabela_padrao(), viagem="T20261027001", arquivo="exemplo.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(x))
    assert wb.sheetnames == ["Painel 面板", "Erradas 错误", "Por rota 线路", "Por segmento 段码", "Turno e dia 班次日期", "Pendências 待确认",
                             "Regras 路由表"]
    erradas = wb["Erradas 错误"]
    assert erradas.max_row == 2 + r.incorretas  # título, cabeçalho e só as erradas (as certas não vão)
    remessas = {erradas.cell(i, 1).value for i in range(3, erradas.max_row + 1)}
    assert remessas == set(res.loc[res["situacao"] == "Incorreta", "remessa"]) and all(isinstance(w, str) for w in remessas)
    rota = wb["Por rota 线路"]
    assert rota.cell(5, 4).number_format == "0.00%" and isinstance(rota.cell(5, 4).value, float)
    assert len(wb["Painel 面板"]._charts) == 4
    # Sem nenhuma errada também gera.
    ok = res[res["situacao"] != "Incorreta"]
    relatorio.gerar_excel(ok, analise.resumir(ok), tabela_padrao())


def test_app_no_modo_demonstracao():
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(PASTA / "app.py"), default_timeout=120)
    at.run()
    assert not at.exception
    at.toggle(key="demo").set_value(True).run()
    botao = next(b for b in at.button if "Buscar no JMS" in b.label)
    botao.click().run()
    assert not at.exception
    assert any("Pronto" in s.value for s in at.success)
    rotulos = [t.label for t in at.tabs]
    assert any("Remessas erradas (" in t for t in rotulos) and any("Pendências (" in t for t in rotulos)
    assert len(at.get("plotly_chart")) == 5
