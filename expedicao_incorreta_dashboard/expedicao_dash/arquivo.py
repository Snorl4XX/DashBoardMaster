"""Abre planilhas pelo CONTEÚDO, não pela extensão: o JMS exporta .xlsx com o nome .xls (ex.: DBGX26102800157.xls), e outros
sistemas salvam página HTML, XML do Excel 2003 ou texto com tabulação como ".xls".

grades(dados) → [(aba, linhas)], cada linha uma lista de valores (texto, número ou datetime), como estão nas células.
"""

from __future__ import annotations

import csv
import io
import re
import xml.etree.ElementTree as ET
from html.parser import HTMLParser

XLSX, XLS, HTML, XML2003, TEXTO = "xlsx", "xls", "html", "xml2003", "texto"


def formato(dados: bytes) -> str:
    if dados[:4] == b"PK\x03\x04":
        return XLSX
    if dados[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return XLS
    if dados[:2] in (b"\xff\xfe", b"\xfe\xff"):  # texto em UTF-16 (alguns sistemas salvam assim)
        cabeca = dados[:8192].decode("utf-16", "ignore").encode("utf-8")
    else:
        cabeca = dados[:4096]
    cabeca = cabeca.lstrip(b"\xef\xbb\xbf \r\n\t").lower()
    if cabeca.startswith(b"<"):
        if b"urn:schemas-microsoft-com:office:spreadsheet" in cabeca and b"<html" not in cabeca:
            return XML2003
        return HTML
    return TEXTO


def grades(dados: bytes, mesclas: bool = False) -> list[tuple[str, list[list]]]:
    """mesclas=True: o valor da célula mesclada vale para todas as células da mescla (planilha de rotas)."""
    if not dados:
        raise ValueError("o arquivo está vazio.")
    f = formato(dados)
    if f == XLSX:
        return _xlsx(dados, mesclas)
    if f == XLS:
        return _xls(dados, mesclas)
    if f == HTML:
        return _html(dados)
    if f == XML2003:
        return _xml2003(dados)
    return [("CSV", _texto(dados))]


def _xlsx(dados: bytes, mesclas: bool) -> list[tuple[str, list[list]]]:
    import openpyxl
    livro = openpyxl.load_workbook(io.BytesIO(dados), data_only=True, read_only=not mesclas)
    out = []
    for ws in livro.worksheets:
        if not mesclas:
            ws.reset_dimensions()  # arquivos exportados por sistemas às vezes dizem o tamanho errado (leria só a 1ª célula)
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        if mesclas:
            for faixa in ws.merged_cells.ranges:
                _preencher(linhas, faixa.min_row - 1, faixa.max_row, faixa.min_col - 1, faixa.max_col,
                           ws.cell(faixa.min_row, faixa.min_col).value)
        out.append((ws.title, linhas))
    livro.close()
    return out


def _xls(dados: bytes, mesclas: bool) -> list[tuple[str, list[list]]]:
    try:
        import xlrd
    except ImportError as e:
        raise ValueError("para abrir .xls antigo falta o pacote xlrd (rode o iniciar.py de novo).") from e
    try:
        livro = xlrd.open_workbook(file_contents=dados, formatting_info=mesclas)
    except NotImplementedError:
        livro = xlrd.open_workbook(file_contents=dados)
    out = []
    for sh in livro.sheets():
        linhas = []
        for i in range(sh.nrows):
            linha = []
            for c in sh.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    try:
                        linha.append(xlrd.xldate.xldate_as_datetime(c.value, livro.datemode))
                    except (ValueError, OverflowError):
                        linha.append(c.value)
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    linha.append(None)
                else:
                    linha.append(c.value)
            linhas.append(linha)
        for rlo, rhi, clo, chi in getattr(sh, "merged_cells", []) or []:
            if rlo < len(linhas) and clo < len(linhas[rlo]):
                _preencher(linhas, rlo, rhi, clo, chi, linhas[rlo][clo])
        out.append((sh.name, linhas))
    return out


def _preencher(linhas: list[list], r0: int, r1: int, c0: int, c1: int, valor) -> None:
    """Linhas r0..r1-1 e colunas c0..c1-1 (índices a partir de 0) recebem o valor."""
    for r in range(r0, r1):
        while len(linhas) <= r:
            linhas.append([])
        linha = linhas[r]
        while len(linha) < c1:
            linha.append(None)
        for c in range(c0, c1):
            linha[c] = valor


def _decodificar(dados: bytes) -> str:
    if dados[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return dados.decode("utf-16")
    for cod in ("utf-8-sig", "cp1252"):
        try:
            return dados.decode(cod)
        except UnicodeDecodeError:
            continue
    return dados.decode("latin-1")


def _texto(dados: bytes) -> list[list]:
    """CSV ou texto com tabulação. O separador é o que mais aparece (o Sniffer erra quando há um título em cima da tabela)."""
    texto = _decodificar(dados)
    sep = max(";,\t|", key=texto[:20000].count)
    return [[v if v != "" else None for v in l] for l in csv.reader(io.StringIO(texto), delimiter=sep)]


class _Tabelas(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tabelas: list[list[list]] = []
        self.linha: list | None = None
        self.celula: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tabelas.append([])
        elif tag == "tr" and self.tabelas:
            self.linha = []
        elif tag in ("td", "th") and self.linha is not None:
            self.celula = []
        elif tag == "br" and self.celula is not None:
            self.celula.append("\n")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.celula is not None and self.linha is not None:
            t = re.sub(r"[ \t\r\f\v]+", " ", "".join(self.celula)).strip()
            self.linha.append(t or None)
            self.celula = None
        elif tag == "tr" and self.linha is not None:
            self.tabelas[-1].append(self.linha)
            self.linha = None

    def handle_data(self, data):
        if self.celula is not None:
            self.celula.append(data)


def _html(dados: bytes) -> list[tuple[str, list[list]]]:
    p = _Tabelas()
    p.feed(_decodificar(dados))
    tabelas = [t for t in p.tabelas if t]
    if not tabelas:
        raise ValueError("o arquivo é uma página HTML sem tabela.")
    return [(f"Tabela {i + 1}", t) for i, t in enumerate(tabelas)]


def _xml2003(dados: bytes) -> list[tuple[str, list[list]]]:
    ns = "{urn:schemas-microsoft-com:office:spreadsheet}"
    raiz = ET.fromstring(_decodificar(dados).encode("utf-8").replace(b'encoding="UTF-16"', b'encoding="UTF-8"'))
    out = []
    for i, ws in enumerate(raiz.iter(f"{ns}Worksheet")):
        linhas = []
        for row in ws.iter(f"{ns}Row"):
            linha: list = []
            for cel in row.findall(f"{ns}Cell"):
                pos = cel.get(f"{ns}Index")
                if pos:
                    while len(linha) < int(pos) - 1:
                        linha.append(None)
                d = cel.find(f"{ns}Data")
                linha.append(d.text if d is not None else None)
            linhas.append(linha)
        out.append((ws.get(f"{ns}Name") or f"Aba {i + 1}", linhas))
    return out
