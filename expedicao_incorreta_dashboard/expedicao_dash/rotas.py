"""Planilha de rotas (as tabelas vermelhas das docas) → regras para saber a rota certa de cada remessa.

Como a operação confere (documento "ferramenta de EXPEDIÇÃO INCORRETA"):
  - ROTA TRONCAL (干线): vale só o Código de primeiro segmento (a sigla). Ex.: AC → SP BRE.
  - ROTA SECUNDÁRIA (支线): vale o par 1º + 2º segmento. Ex.: GRU 404-05 → F S-PENHA-SP.
  - DC TBT-SP, DC GRU-SP e SBN estão na tabela troncal, mas seguem o padrão da secundária (têm 2º Nível): vale o par.
  - Linha sem 2º Nível (ex.: DC S-VLANA-SP · VLANA; DC NOD-SP · NOD) vale pelo 1º segmento.
  Para cada remessa: primeiro o par exato (1º + 2º); se nenhuma linha tem o par, o 1º segmento sozinho.

Leitura da planilha:
  - Acha sozinha a linha do cabeçalho e as colunas ("Rota de Carregamento Após o Ajuste", "1º Nível", "2º Nível", "Novo número
    da doca", "Observações") em cada aba. Também lê a planilha simples (Tipo, Rota, 1º Nível, 2º Nível, Observação).
  - Células mescladas (a mesma rota em duas docas, ex.: DC SSZ-SP 08 e 09) valem para todas as linhas.
  - Troncal ou secundária: pela aba ou pelo título (支线 / Regionais / Secundária = secundária; 干线 / Linha / Troncal = troncal).
Sem planilha nenhuma, vale a tabela de 24/10 embutida aqui (ROTAS_24_10).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from .texto import chave_rota, limpo, lista_primeiros, lista_segundos, norm, norm_primeiro, norm_segundo, partes_rota

TRONCAL, SECUNDARIA = "Troncal", "Secundária"
# Pedaços que o JMS põe a mais no nome da rota sem mudar qual é ela (tipo de unidade e UF): "PI THE" = THE; "DC SBN-SP" = SBN.
ENFEITES = frozenset({"DC", "CD", "SC", "F", "HUB", "BASE"} | set(
    "AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SP SE TO".split()))


@dataclass(frozen=True)
class Regra:
    rota: str
    tipo: str
    primeiros: tuple[str, ...]
    segundos: tuple[str, ...] = ()  # vazio = vale só o 1º segmento
    docas: str = ""
    obs: str = ""
    origem: str = ""  # aba e linha da planilha

    @property
    def por_par(self) -> bool:
        return bool(self.segundos)


@dataclass
class Esperada:
    rota: str = ""  # a rota certa ("" = não deu para saber)
    situacao: str = ""  # ok · sem_codigo · segmento_fora · codigo_fora · ambigua
    metodo: str = ""  # "1º + 2º segmento" ou "1º segmento"
    regra: Regra | None = None
    candidatas: tuple[str, ...] = ()


@dataclass
class TabelaRotas:
    regras: list[Regra]
    nome: str = "Planilha de rotas"
    avisos: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.pares: dict[tuple[str, str], dict[str, Regra]] = {}
        self.prefixos: dict[str, dict[str, Regra]] = {}
        self.com_par: set[str] = set()  # 1º segmento que aparece em alguma linha com 2º Nível (ex.: GRU)
        self.tipo_da_rota: dict[str, str] = {}
        for r in self.regras:
            self.tipo_da_rota.setdefault(r.rota, r.tipo)
            for p in r.primeiros:
                for chave in _apelidos_codigo(p):
                    if r.segundos:
                        self.com_par.add(chave)
                        for s in r.segundos:
                            self.pares.setdefault((chave, s), {})[r.rota] = r
                    else:
                        self.prefixos.setdefault(chave, {})[r.rota] = r
        self.rotas = sorted(self.tipo_da_rota, key=norm)
        self._por_chave = {chave_rota(r): r for r in self.rotas}
        self._partes = {r: partes_rota(r) for r in self.rotas}
        for (p, s), rotas in self.pares.items():
            if len(rotas) > 1:
                self.avisos.append(f"O segmento {p} {s} está em mais de uma rota: {', '.join(sorted(rotas))}.")
        for p, rotas in self.prefixos.items():
            if len(rotas) > 1:
                self.avisos.append(f"O código {p} está em mais de uma rota: {', '.join(sorted(rotas))}.")

    # ---------------------------------------------------------------- rota certa
    def esperada(self, primeiro: str, segundo: str = "") -> Esperada:
        p, s = norm_primeiro(primeiro), norm_segundo(segundo)
        if not p:
            return Esperada(situacao="sem_codigo")
        achadas = self.pares.get((p, s), {}) if s else {}
        metodo = "1º + 2º segmento"
        if not achadas:
            achadas, metodo = self.prefixos.get(p, {}), "1º segmento"
        if len(achadas) == 1:
            rota, regra = next(iter(achadas.items()))
            return Esperada(rota, "ok", metodo, regra)
        if len(achadas) > 1:
            return Esperada(situacao="ambigua", candidatas=tuple(sorted(achadas)))
        return Esperada(situacao="segmento_fora" if p in self.com_par else "codigo_fora")

    # ---------------------------------------------------------------- nome da rota carregada
    def resolver(self, nome: str) -> str:
        """Nome da rota carregada (ex.: a "Próxima parada" do JMS) → rota da planilha. "" = não reconhecida.
        Mesmo nome sem contar espaços e traços ("DC S-MOO 02-SP" = "DC S-MOO02-SP"); ou, se o nome não é de nenhuma rota, a única
        rota cujos pedaços estão todos nele: "PI THE" → THE; "DC SBN-SP" → SBN."""
        k = chave_rota(nome)
        if not k:
            return ""
        if k in self._por_chave:
            return self._por_chave[k]
        pedacos = partes_rota(nome)
        # Sobra só "enfeite" (DC, F, a UF…): "F S-PENHA 04-SP" NÃO vira "F S-PENHA-SP" (o 04 faz dela outra rota).
        achadas = [r for r, ps in self._partes.items() if ps and ps <= pedacos and (pedacos - ps) <= ENFEITES]
        if len(achadas) > 1:  # fica a mais completa, se só uma for
            maior = max(len(self._partes[r]) for r in achadas)
            achadas = [r for r in achadas if len(self._partes[r]) == maior]
        return achadas[0] if len(achadas) == 1 else ""

    def linhas(self) -> list[dict]:
        """As regras como a tela e o Excel mostram."""
        return [{"Rota": r.rota, "Tipo": r.tipo, "1º Nível": ", ".join(r.primeiros), "2º Nível": ", ".join(r.segundos),
                 "Confere por": "1º + 2º segmento" if r.segundos else "1º segmento", "Docas": r.docas, "Observações": r.obs,
                 "Origem": r.origem} for r in self.regras]


def _apelidos_codigo(p: str) -> list[str]:
    """"DC-BYE" (como está na tabela "PE、DC—BYE") também vale como "BYE" (como o JMS pode mandar)."""
    out = [p]
    m = re.fullmatch(r"DC-?([A-Z0-9]{2,})", p)
    if m:
        out.append(m.group(1))
    return out


# ------------------------------------------------------------------ ler a planilha
_RE_ROTA = re.compile(r"rota de carregamento|装车线路|^rota\b|proxima parada|下一站|^linha$")
_RE_P1 = re.compile(r"1\s*o?\s*nivel|一段码|primeiro segmento")
_RE_P2 = re.compile(r"2\s*o?\s*nivel|二段码|dois segmentos|segundo segmento")
_RE_DOCA = re.compile(r"novo numero da doca|新码头|^doca")
_RE_OBS = re.compile(r"observa|备注")
_RE_TIPO = re.compile(r"^tipo")


def _tipo_do_texto(texto: str) -> str:
    t = norm(texto)
    if re.search(r"支线|regiona|secund", t):
        return SECUNDARIA
    if re.search(r"干线|troncal|\blinha\b", t):
        return TRONCAL
    return ""


def _grade_openpyxl(ws) -> list[list]:
    """Valores da aba com as células mescladas preenchidas (a rota mesclada em duas docas vale para as duas)."""
    linhas = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    for faixa in ws.merged_cells.ranges:
        valor = ws.cell(faixa.min_row, faixa.min_col).value
        for r in range(faixa.min_row, faixa.max_row + 1):
            while len(linhas) < r:
                linhas.append([])
            linha = linhas[r - 1]
            while len(linha) < faixa.max_col:
                linha.append(None)
            for c in range(faixa.min_col, faixa.max_col + 1):
                linha[c - 1] = valor
    return linhas


def _abas(dados: bytes, nome_arquivo: str) -> list[tuple[str, list[list]]]:
    if nome_arquivo.lower().endswith(".xls"):
        import pandas as pd  # .xls (antigo): sem células mescladas; a rota vazia herda a de cima (mesma doca dupla)
        livro = pd.read_excel(io.BytesIO(dados), sheet_name=None, header=None, dtype=object, engine="xlrd")
        return [(nome, df.astype(object).where(df.notna(), None).values.tolist()) for nome, df in livro.items()]
    import openpyxl
    livro = openpyxl.load_workbook(io.BytesIO(dados), data_only=True)
    return [(ws.title, _grade_openpyxl(ws)) for ws in livro.worksheets]


def _cabecalho(grade: list[list]) -> tuple[int, dict[str, int]] | None:
    for i, linha in enumerate(grade[:20]):
        cols: dict[str, int] = {}
        for j, v in enumerate(linha):
            t = norm(v)
            if not t:
                continue
            for chave, rx in (("p1", _RE_P1), ("p2", _RE_P2), ("rota", _RE_ROTA), ("doca", _RE_DOCA), ("obs", _RE_OBS), ("tipo", _RE_TIPO)):
                if chave not in cols and rx.search(t):
                    cols[chave] = j
                    break
        if "p1" in cols and "rota" in cols:
            return i, cols
    return None


def ler_rotas(dados: bytes, nome_arquivo: str = "rotas.xlsx") -> TabelaRotas:
    """Planilha de rotas (.xlsx ou .xls) → TabelaRotas. Erro (ValueError) se nenhuma aba tiver as colunas Rota e 1º Nível."""
    regras: list[Regra] = []
    docas_da_rota: dict[str, list[str]] = {}
    avisos: list[str] = []
    for nome_aba, grade in _abas(dados, nome_arquivo):
        achado = _cabecalho(grade)
        if not achado:
            continue
        i_cab, cols = achado
        tipo_aba = _tipo_do_texto(nome_aba) or _tipo_do_texto(" ".join(limpo(v) for linha in grade[:i_cab + 1] for v in linha))
        herdada = ""
        for i in range(i_cab + 1, len(grade)):
            linha = grade[i]
            cel = lambda k: limpo(linha[cols[k]]) if k in cols and cols[k] < len(linha) else ""  # noqa: E731
            rota_txt = re.sub(r"\s+", " ", cel("rota"))
            if nome_arquivo.lower().endswith(".xls") and not rota_txt and (cel("p1") or cel("doca")):
                rota_txt = herdada
            herdada = rota_txt if rota_txt not in ("-", "—") else ""
            if not rota_txt or rota_txt in ("-", "—"):
                continue
            doca = cel("doca")
            if doca and doca not in ("-", "—"):
                docas_da_rota.setdefault(rota_txt, [])
                if doca not in docas_da_rota[rota_txt]:
                    docas_da_rota[rota_txt].append(doca)
            primeiros = lista_primeiros(cel("p1"))
            if not primeiros:
                continue
            segundos = lista_segundos(cel("p2"))
            tipo = _tipo_do_texto(cel("tipo")) or (SECUNDARIA if segundos else tipo_aba) or TRONCAL
            regras.append(Regra(rota_txt, tipo, tuple(primeiros), tuple(segundos), "", cel("obs").replace("\n", " · "),
                                f"{nome_aba}, linha {i + 1}"))
    if not regras:
        raise ValueError("não achei as colunas \"Rota de Carregamento\" e \"1º Nível\" em nenhuma aba.")
    # Linhas iguais (doca dupla mesclada) viram uma só; as docas da rota vão juntas.
    unicas: dict[tuple, Regra] = {}
    for r in regras:
        k = (r.rota, r.primeiros, r.segundos)
        if k not in unicas:
            docas = ", ".join(docas_da_rota.get(r.rota, []))
            unicas[k] = Regra(r.rota, r.tipo, r.primeiros, r.segundos, docas, r.obs, r.origem)
    return TabelaRotas(list(unicas.values()), nome_arquivo, avisos)


# ------------------------------------------------------------------ tabela de 24/10 (embutida)
def _seg(frente: str, nums: list[int]) -> tuple[str, ...]:
    return tuple(f"{frente}-{n:02d}" for n in nums)


ROTAS_24_10 = [
    # Troncal (干线)
    Regra("GO GYN", TRONCAL, ("GO",), (), "05"),
    Regra("SC BNU", TRONCAL, ("SC", "SC1"), (), "06"),
    Regra("RS NSR", TRONCAL, ("RS", "RS1"), (), "07"),
    Regra("SBN", SECUNDARIA, ("SBN",), _seg("375", list(range(11))), "08"),
    Regra("CE FOR", TRONCAL, ("CE",), (), "09"),
    Regra("RJ SJM", TRONCAL, ("RJ", "RJ1", "VDC", "CDG"), (), "10", "Adicionar novo código inicial: CDG (DC CDG-RJ)"),
    Regra("SE AJU", TRONCAL, ("SE", "AL"), (), "15"),
    Regra("PR SJS", TRONCAL, ("PR", "PR1", "CAC"), (), "16"),
    Regra("PE JGS", TRONCAL, ("PE", "DC-BYE"), (), "17"),
    Regra("BA FEC", TRONCAL, ("BA", "SBA", "SNS"), (), "18"),
    Regra("MG CGE", TRONCAL, ("MG", "MG1", "CHV"), (), "19"),
    Regra("PA ANA", TRONCAL, ("PA",), (), "20", "Movimentação manual de sacas"),
    Regra("SP BRE", TRONCAL, ("AC", "AM", "BAU", "BVB", "JDF", "LDB", "MIA", "MRB", "RO", "SJP", "SOD", "SP", "SP1", "STM", "TO",
                              "VCP", "XAP", "BJE", "MA", "PPB", "NAT"), (), "21, 22", "Excluir CDG"),
    Regra("DC TBT-SP", SECUNDARIA, ("SP",), _seg("402", [1, 3, 5]) + ("421-00",), "23", "SP (antigo) · transferido da doca 14 para a 23"),
    Regra("DC TBT-SP", SECUNDARIA, ("GRU",), _seg("402", [0, 1, 3, 5, 6, 7, 8, 2, 4]), "23", "GRU (novo)"),
    Regra("DC GRU-SP", SECUNDARIA, ("SP",), _seg("395", [5, 6, 7]) + ("381-04",), "44", "SP (antigo)"),
    Regra("DC GRU-SP", SECUNDARIA, ("GRU",), _seg("381", [0, 1, 2, 3, 4, 5, 11, 13, 14, 15, 18, 19, 20, 21]) + ("378-04",), "44",
          "GRU (novo) · 378-04: data de ativação a definir"),
    Regra("MS CGR", TRONCAL, ("MS",), (), "43"),
    Regra("MT CGB", TRONCAL, ("MT", "JPN"), (), "42"),
    Regra("THE", TRONCAL, ("PI", "IMP"), (), "41"),
    Regra("ES SRR", TRONCAL, ("ES",), (), "40"),
    Regra("SP RBP", TRONCAL, ("NE", "NE1"), (), "39"),
    Regra("DF BSB", TRONCAL, ("DF",), (), "38"),
    # Secundária (支线)
    Regra("DC S-VLANA-SP", SECUNDARIA, ("VLANA",), (), "01", "Reservado"),
    Regra("F S-JBO-SP", SECUNDARIA, ("GRU",), ("404-11",), "02", "Reservado"),
    Regra("F S-SAP 02-SP", SECUNDARIA, ("GRU",), ("404-12",), "03", "Reservado"),
    Regra("F S-MGUE 02-SP", SECUNDARIA, ("GRU",), ("404-08",), "04", "Com coluna de sustentação"),
    Regra("F S-PQBOTU-SP", SECUNDARIA, ("GRU",), ("404-09",), "05"),
    Regra("F VLMT-SP", SECUNDARIA, ("GRU",), ("404-10",), "06"),
    Regra("S-MGUE-SP", SECUNDARIA, ("GRU",), ("404-01",), "07"),
    Regra("DC SSZ-SP", SECUNDARIA, ("GRU",), _seg("378", [0, 1, 2, 3, 5, 6]), "08, 09", "Com coluna de sustentação"),
    Regra("S-MGUE 02-SP", SECUNDARIA, ("GRU",), ("404-13",), "10", "Reservado"),
    Regra("DC S-MOO 02-SP", SECUNDARIA, ("SP",), ("364-00", "435-00", "350-00", "365-00", "381-12") + _seg("430", [1, 5, 7, 8, 9]),
          "11, 12", "SP (antigo)"),
    Regra("DC S-MOO 02-SP", SECUNDARIA, ("MOO",), _seg("430", list(range(10))), "11, 12", "MOO (novo) · reajustado"),
    Regra("S-PENHA-SP", SECUNDARIA, ("GRU",), ("404-02",), "13"),
    Regra("F S-PENHA 02-SP", SECUNDARIA, ("GRU",), ("404-07",), "14", "Com coluna de sustentação"),
    Regra("S-SRAF-SP", SECUNDARIA, ("GRU",), ("404-00",), "15, 16"),
    Regra("DC NOD-SP", SECUNDARIA, ("NOD",), (), "17, 18", "Reservado"),
    Regra("S-SAPOP-SP", SECUNDARIA, ("GRU",), ("404-03",), "19"),
    Regra("F S-MGUE-SP", SECUNDARIA, ("GRU",), ("404-04",), "20", "Com coluna de sustentação"),
    Regra("F S-PENHA-SP", SECUNDARIA, ("GRU",), ("404-05",), "22"),
    Regra("F S-PENHA 03-SP", SECUNDARIA, ("GRU",), ("404-06",), "24"),
]


def tabela_padrao() -> TabelaRotas:
    return TabelaRotas([Regra(r.rota, r.tipo, r.primeiros, r.segundos, r.docas, r.obs, "tabela de 24/10 (embutida)")
                        for r in ROTAS_24_10], "Tabela de 24/10 (embutida no app)")
