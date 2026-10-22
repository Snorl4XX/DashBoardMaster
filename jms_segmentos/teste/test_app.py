"""Teste do app JMS_Segmentos corrigido contra o JMS falso (tudo fictício).

Uso: python -I test_app.py <pasta do app> <pasta de saída>
"""
import asyncio
import base64
import io
import json
import os
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
from pathlib import Path

APP_DIR, OUT = Path(sys.argv[1]), Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(1, str(APP_DIR))
os.environ["JMS_SEGMENTOS_HEADLESS"] = "1"

import uvicorn  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

import jms_falso  # noqa: E402
import app as m  # noqa: E402

PJMS, PAPP = 8791, 8792
falhas = 0


def check(ok, nome, extra=""):
    global falhas
    print(("OK   " if ok else "FALHOU ") + nome + (f"  -> {extra}" if (extra and not ok) else ""))
    if not ok:
        falhas += 1


def api(path, body=None, timeout=600):
    data = None if body is None else json.dumps(body).encode()
    req = urllib.request.Request(f"http://127.0.0.1:{PAPP}{path}", data=data, headers={"Content-Type": "application/json"}, method="GET" if body is None else "POST")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            raw = r.read()
            return r.status, (json.loads(raw) if r.headers.get("Content-Type", "").startswith("application/json") else raw)
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read() or b"{}")


def esperar(cond, segundos=30):
    fim = time.time() + segundos
    while time.time() < fim:
        if cond():
            return True
        time.sleep(0.2)
    return False


def no_navegador_do_app(coro_fn):
    return asyncio.run_coroutine_threadsafe(coro_fn(), m._pw_loop).result(60)


def conferir(rows, nome):
    errados, transitorias = [], 0
    for r in rows:
        exp = jms_falso.esperado(r["remessa"])
        if r.get("sessao"):
            errados.append((r["remessa"], "sessao"))
        elif exp is None:
            if not r.get("nao_encontrada"):
                errados.append((r["remessa"], r))
        elif r.get("erro"):
            transitorias += 1
            errados.append((r["remessa"], r.get("erro")))
        elif (r.get("codigo_primeiro_segmento", ""), r.get("codigo_dois_segmentos", "")) != exp:
            errados.append((r["remessa"], r))
    check(not errados, f"{nome}: todas as remessas com os 2 códigos certos (3º código fora)", errados[:5])
    return transitorias


fake, est = jms_falso.criar(PJMS)
m.JMS_URL = f"http://127.0.0.1:{PJMS}/app/operatingplatform/index#tab=1"
m.PROFILE_DIR = Path(tempfile.mkdtemp(prefix="perfil_"))
srv = uvicorn.Server(uvicorn.Config(m.app, host="127.0.0.1", port=PAPP, log_level="warning"))
threading.Thread(target=srv.run, daemon=True).start()
check(esperar(lambda: srv.started, 15), "servidor do app subiu")

# 1. Abre o Chromium (antes: "Playwright Sync API inside the asyncio loop") e captura a consulta feita em OUTRA aba.
st, r = api("/api/browser/open", {})
check(st == 200 and r.get("ok"), "abre o Chromium", r)
check(esperar(lambda: api("/api/status")[1].get("captured")), "captura a consulta de teste feita na aba do Rastreamento")
st, s = api("/api/status")
check(s.get("seed_waybill") == "777000000000001" and s.get("browser_open"), "status: navegador aberto e remessa de teste", s)
hdrs = m._capture["headers"]
check(not any(k.startswith(":") for k in hdrs) and "cookie" not in hdrs and hdrs.get("authtoken") == "tok-ficticio-1", "cabeçalhos capturados sem :pseudo e sem cookie", list(hdrs))

# 2. Planilha grande: 3.000 remessas numa aba com cabeçalho, outra aba sem nada a ver, repetidas.
N = 3000
wb = Workbook()
wb.active.title = "Outra"
wb.active.append(["nada", "aqui"])
ws = wb.create_sheet("Base")
ws.append(["Data", "Remessa", "Cliente"])
remessas = [f"777{i:012d}" for i in range(1, N + 1)]
for i, w in enumerate(remessas):
    ws.append(["2026-10-22", int(w), "Cliente Ficticio"])
for w in remessas[:40]:
    ws.append(["2026-10-22", w, "Repetida"])
buf = io.BytesIO()
wb.save(buf)
xlsx_path = OUT / "planilha_teste.xlsx"
xlsx_path.write_bytes(buf.getvalue())
st, up = api("/api/sheets/inspect", {"filename": "teste.xlsx", "content_base64": base64.b64encode(buf.getvalue()).decode()})
base = next(x for x in up["sheets"] if x["name"] == "Base")
check(st == 200 and base["suggested_index"] == 1 and base["header_row"] == 1, "acha a aba e sugere a coluna Remessa", up["sheets"])
st, lido = api("/api/sheets/read", {"upload_id": up["upload_id"], "sheet_name": "Base", "column_index": 1})
check(lido["count"] == N and lido["waybills"][:2] == remessas[:2], "lê as 3.000 remessas (repetidas contam uma vez)", lido.get("count"))

# 3. JMS cheio: 25% das consultas falham (502, "系统繁忙", vazio, sem dados). Blocos de 1.000 (como o site externo manda).
est.falha = 0.25
t0 = time.time()
todas = []
for i in range(0, N, 1000):
    st, part = api("/api/query", {"waybills": lido["waybills"][i:i + 1000], "patterns": []})
    check(st == 200 and len(part["_raw_rows"]) == len(lido["waybills"][i:i + 1000]) and not part["sessao_expirada"], f"bloco {i // 1000 + 1}: respondeu todas", (st, part.get("detail")))
    todas += part["_raw_rows"]
dt = time.time() - t0
# Igual ao botão "Consultar as que faltaram": de novo só as que ficaram com erro passageiro.
faltaram = [r["remessa"] for r in todas if r.get("erro") and not r.get("nao_encontrada")]
print(f"     {N} remessas em {dt:.1f}s · {len(faltaram)} ficaram com erro depois de 6 tentativas (JMS com 25% de falhas)")
check(len(faltaram) <= 5, "quase todas resolvidas na 1ª rodada", len(faltaram))
if faltaram:
    st, extra = api("/api/query", {"waybills": faltaram})
    novas = {r["remessa"]: r for r in extra["_raw_rows"]}
    todas = [novas.get(r["remessa"], r) for r in todas]
trans = conferir(todas, "JMS com 25% de falhas")
print(f"     {est.consultas} consultas ao JMS · no máximo {est.pico} ao mesmo tempo")
check(est.pico <= m.PARALLEL, "não passa de 8 consultas ao mesmo tempo", est.pico)

# 4. Sessão vence no meio: o bloco para, devolve o que pegou e marca o resto; nova captura e "as que faltaram" completa.
est.falha = 0.0
lote = [f"777{i:012d}" for i in range(5001, 5601)]
est.expira_em = est.consultas + 200
st, part = api("/api/query", {"waybills": lote})
rows = part["_raw_rows"]
pegas = [r for r in rows if not r.get("sessao")]
check(part["sessao_expirada"] and 150 <= len(pegas) < 600 and all("Sessão do JMS expirada" in r["erro"] for r in rows if r.get("sessao")), "sessão vencida: para e avisa (sem gastar o resto)", (part.get("sessao_expirada"), len(pegas)))
est.token = "tok-ficticio-2"


async def nova_consulta():
    page = m._context.pages[-1]
    await page.goto(f"http://127.0.0.1:{PJMS}/app/tracking?nova=1#auto=1&tok=tok-ficticio-2")


no_navegador_do_app(nova_consulta)
check(esperar(lambda: m._capture and m._capture["headers"].get("authtoken") == "tok-ficticio-2"), "nova consulta de teste atualiza a captura")
faltam = [r["remessa"] for r in rows if r.get("sessao")]
st, part2 = api("/api/query", {"waybills": faltam})
juntas = {r["remessa"]: r for r in pegas}
juntas.update({r["remessa"]: r for r in part2["_raw_rows"]})
conferir([juntas[w] for w in lote], "depois de entrar de novo")

# 5. Resultado, gráfico e Excel.
st, res = api("/api/aggregate", {"rows": todas})
nao = sum(1 for w in remessas if jms_falso.esperado(w) is None)
sem = sum(1 for w in remessas if jms_falso.esperado(w) == ("", ""))
check(res["remessas_consultadas"] == N and res["remessas_nao_encontradas"] == nao and res["remessas_sem_codigo"] == sem and res["remessas_com_dados"] == N - nao - sem,
      "contagens: com código / sem código / não encontradas", {k: res[k] for k in res if k.startswith("remessas")})
g1 = {x["codigo"]: x["quantidade"] for x in res["grafico"] if x["campo"] == m.FIELD_1}
check(set(g1) == set(jms_falso.S1) and sum(g1.values()) == N - nao - sem, "gráfico do 1º código soma as remessas com código", g1)
check(not any(x["codigo"] == "900" for x in res["grafico"]), "3º código (900) fica fora")
st, xbytes = api("/api/export", {"result": res})
xl = load_workbook(io.BytesIO(xbytes))
r1 = list(xl["Resultado"].iter_rows(values_only=True))
check(xl.sheetnames[:2] == ["Resultado", "Resumo"] and r1[0] == ("Remessa", m.FIELD_1, m.FIELD_2) and len(r1) == N + 1, "Excel: Resultado com remessa e os 2 códigos", (xl.sheetnames, r1[0], len(r1)))
check(len(xl["Resumo"]._charts) == 2 and "Sem código" in xl.sheetnames and xl["Sem código"].max_row == nao + sem + 1, "Excel: Resumo com 2 gráficos e aba Sem código", (len(xl["Resumo"]._charts), xl.sheetnames))
(OUT / "resultado_api.xlsx").write_bytes(xbytes)

# 6. Tela do app (http://127.0.0.1:8765 no uso real): planilha → aba → coluna → consultar → gráficos, tabela, Excel.
est.falha = 0.15
from playwright.sync_api import sync_playwright  # noqa: E402

with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={"width": 1366, "height": 900}, accept_downloads=True)
    erros_js = []
    pg.on("pageerror", lambda e: erros_js.append(str(e)))
    pg.goto(f"http://127.0.0.1:{PAPP}/")
    pg.wait_for_function("document.getElementById('capturePill').textContent.includes('capturada')", timeout=15000)
    pg.set_input_files("#file", str(xlsx_path))
    pg.wait_for_function("document.getElementById('sheetMsg').textContent.includes('remessas na aba')", timeout=20000)
    check("3.000 remessas" in pg.inner_text("#sheetMsg") and pg.eval_on_selector("#aba", "e => e.selectedOptions[0].text") == "Base", "tela: escolhe a aba Base e lê 3.000", pg.inner_text("#sheetMsg"))
    t0 = time.time()
    pg.click("#queryBtn")
    pg.wait_for_function("document.getElementById('queryMsg').textContent.includes('Pesquisa concluída')", timeout=600000)
    print(f"     tela: {N} remessas em {time.time() - t0:.1f}s")
    check(pg.inner_text("#mTotal") == "3.000" and pg.inner_text("#mOk") == f"{N - nao - sem:,}".replace(",", "."), "tela: cartões", [pg.inner_text(x) for x in ("#mTotal", "#mOk", "#mSem", "#mErro")])
    check(pg.locator("#chart1 .barrow").count() == 4 and pg.locator("#chart2 .barrow").count() == 15, "tela: os 2 gráficos", (pg.locator("#chart1 .barrow").count(), pg.locator("#chart2 .barrow").count()))
    check(pg.locator("#tab-rows tbody tr").count() == 1000 and "Mostrando 1.000 de 3.000" in pg.inner_text("#rowsInfo"), "tela: tabela de remessas", pg.inner_text("#rowsInfo"))
    pg.fill("#search", remessas[1])
    linha = pg.locator("#tab-rows tbody tr").first.inner_text().split("\t")
    check(linha[:3] == [remessas[1], *jms_falso.esperado(remessas[1])] and linha[3] == "OK", "tela: busca acha a remessa com os 2 códigos", linha)
    pg.fill("#search", "")
    pg.screenshot(path=str(OUT / "tela_resultado.png"), full_page=True)
    with pg.expect_download() as d:
        pg.click("text=Baixar planilha Excel")
    d.value.save_as(str(OUT / "resultado_tela.xlsx"))
    xl2 = load_workbook(OUT / "resultado_tela.xlsx")
    check(xl2["Resultado"].max_row == N + 1, "tela: baixa o Excel com todas as remessas")
    check(not erros_js, "tela: sem erro de JavaScript", erros_js)
    pg.set_viewport_size({"width": 390, "height": 800})
    pg.screenshot(path=str(OUT / "tela_celular.png"), full_page=False)
    b.close()

st, r = api("/api/browser/close", {})
check(st == 200 and not api("/api/status")[1]["browser_open"], "fecha o navegador")
print("\nRESULTADO:", "tudo certo" if not falhas else f"{falhas} falha(s)")
srv.should_exit = True
sys.exit(1 if falhas else 0)
