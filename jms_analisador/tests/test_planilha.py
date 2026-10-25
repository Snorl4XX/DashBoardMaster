import io

from openpyxl import load_workbook

from jmslib import config
from jmslib.demo import JMSFicticio, remessas_demo
from jmslib.execucao import Opcoes, executar, gerar_planilha


def _resultado():
    op = Opcoes(padroes=["O expresso chegou [SP GRU] A parada anterior foi", config.PADRAO_EXEMPLO], segmentos=True, avaria=True,
                modelo_imagem=JMSFicticio.MODELO_IMAGEM)
    return executar(JMSFicticio(), remessas_demo(30), op)


def test_planilha_padrao_jt():
    res = _resultado()
    wb = load_workbook(io.BytesIO(gerar_planilha(res, None, True)))
    assert wb.sheetnames[:4] == ["Resumo 汇总", "Pesquisa 搜索明细", "Segmentos 三段码", "Avaria 破损明细"]
    ws = wb["Pesquisa 搜索明细"]
    # Cabeçalho: português em cima e mandarim embaixo, fundo vermelho.
    assert ws["A1"].value == "Remessa" and ws["A2"].value == "运单号"
    assert ws["A1"].fill.fgColor.rgb.endswith("E60012")
    assert ws["A1"].font.color.rgb.endswith("FFFFFF")
    assert ws.auto_filter.ref.startswith("A2:") and ws.freeze_panes == "B3"
    cabecalhos = [c.value for c in ws[1]]
    for esperado in ("Tempo de digitalização (1º bipe)", "Turno", "Origem", "Quem escaneou (login)", "ID de Viagem",
                     "Número da Saca", "Imagem", "Código de primeiro segmento", "SC Destino"):
        assert esperado in cabecalhos
    assert len(ws._images) > 0  # miniaturas dentro das linhas
    resumo = wb["Resumo 汇总"]
    assert len(resumo._charts) >= 10
    assert len(wb["Avaria 破损明细"]._images) > 0


def test_planilha_sem_miniaturas_e_so_segmentos():
    res = executar(JMSFicticio(), remessas_demo(8), Opcoes(segmentos=True))
    wb = load_workbook(io.BytesIO(gerar_planilha(res, None, False)))
    assert "Segmentos 三段码" in wb.sheetnames and "Pesquisa 搜索明细" not in wb.sheetnames
    ws = wb["Segmentos 三段码"]
    assert [c.value for c in ws[1]][:3] == ["Remessa", "Código de primeiro segmento", "Código de dois segmento"]
    assert [c.value for c in ws[2]][:3] == ["运单号", "一段码", "二段码"]


def test_token_recusado_ainda_gera_planilha_com_erros():
    res = executar(JMSFicticio(token_recusado=True), remessas_demo(5), Opcoes(padroes=["x"]))
    assert res.parado_por
    wb = load_workbook(io.BytesIO(gerar_planilha(res, None, True)))
    assert "Erros 错误" in wb.sheetnames
