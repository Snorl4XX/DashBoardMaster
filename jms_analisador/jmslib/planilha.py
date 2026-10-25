"""Planilha no padrão J&T: cabeçalho vermelho com português em cima e mandarim embaixo, gráficos do Excel e a miniatura da
imagem dentro da célula "Imagem" de cada linha (anda junto ao ordenar e filtrar)."""

from __future__ import annotations

import io
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from typing import Any, Callable

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.drawing.image import Image as ImagemXL
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import config
from .analise import contar, contar_por_hora, segmentos_juntados
from .imagens import bytes_de_data_url, miniatura_png
from .texto import para_data

VERMELHO = config.VERMELHO_JT.lstrip("#")
VERMELHO_ESCURO = "B8000E"
CINZA_LINHA = "E5E7EB"
_FONTE = "Arial"

_CAB_PT = PatternFill("solid", fgColor=VERMELHO)
_CAB_ZH = PatternFill("solid", fgColor=VERMELHO_ESCURO)
_BRANCO = Font(name=_FONTE, color="FFFFFF", bold=True, size=10)
_BRANCO_ZH = Font(name=_FONTE, color="FFFFFF", bold=False, size=9)
_NORMAL = Font(name=_FONTE, size=10)
_LINK = Font(name=_FONTE, size=10, color="0563C1", underline="single")
_BORDA = Border(bottom=Side(style="thin", color=CINZA_LINHA))
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_MEIO = Alignment(vertical="center")

_COLUNAS_DATA = {"tempo_digitalizacao", "recebimento", "tempo_carregamento", "checking", "deslacre", "tempo_referencia",
                 "bipe_avaria"}
_LARGURAS = {"descricao": 70, "padrao": 34, "imagem": 20, "imagem_avaria": 20, "link_imagem": 16, "link_imagem_avaria": 16,
             "erro": 80, "remessa": 18}
_LADO_MINIATURA = 110  # px
_ALTURA_LINHA_FOTO = 90  # pt (120 px)

Progresso = Callable[[int, int, str], None]


# ---------------------------------------------------------------------- miniaturas
def baixar_miniaturas(urls: list[str], baixar: Callable[[str], bytes] | None, paralelo: int = 8,
                      progresso: Progresso | None = None) -> dict[str, tuple[bytes, int, int]]:
    """URL → (png, largura, altura). Foto que não baixa fica de fora (a célula mantém o link)."""
    unicos = [u for u in dict.fromkeys(urls) if u]
    saida: dict[str, tuple[bytes, int, int]] = {}
    aviso = progresso or (lambda a, b, c: None)

    def uma(url: str):
        try:
            dados = bytes_de_data_url(url) if url.startswith("data:") else (baixar(url) if baixar else b"")
            return url, miniatura_png(dados, _LADO_MINIATURA) if dados else None
        except Exception:  # noqa: BLE001 — sem a miniatura, fica o link
            return url, None

    with ThreadPoolExecutor(max_workers=max(1, paralelo)) as pool:
        for n, (url, mini) in enumerate(pool.map(uma, unicos), start=1):
            if mini:
                saida[url] = mini
            aviso(n, len(unicos), "Baixando as imagens para a planilha")
    return saida


def _px_largura(largura_coluna: float) -> int:
    return int(largura_coluna * 7 + 5)


def _ancorar_imagem(ws: Worksheet, png: bytes, w: int, h: int, linha: int, coluna: int, largura_coluna: float) -> None:
    """Miniatura dentro da célula (âncora de duas células na própria célula: move e filtra junto com a linha)."""
    img = ImagemXL(io.BytesIO(png))
    img.width, img.height = w, h
    emu = 9525
    caixa_w = _px_largura(largura_coluna)
    caixa_h = int(_ALTURA_LINHA_FOTO * 96 / 72)
    off_x = max(2, (caixa_w - w) // 2)
    off_y = max(2, (caixa_h - h) // 2)
    inicio = AnchorMarker(col=coluna - 1, colOff=off_x * emu, row=linha - 1, rowOff=off_y * emu)
    fim = AnchorMarker(col=coluna - 1, colOff=min(caixa_w - 1, off_x + w) * emu, row=linha - 1, rowOff=min(caixa_h - 1, off_y + h) * emu)
    img.anchor = TwoCellAnchor(editAs="twoCell", _from=inicio, to=fim)
    ws.add_image(img)


# ---------------------------------------------------------------------- tabelas
def _cabecalho(ws: Worksheet, linha: int, coluna: int, chaves: list[str]) -> None:
    for j, chave in enumerate(chaves):
        pt, zh = config.COLUNAS.get(chave, (chave, chave))
        c1 = ws.cell(row=linha, column=coluna + j, value=pt)
        c2 = ws.cell(row=linha + 1, column=coluna + j, value=zh)
        c1.fill, c1.font, c1.alignment = _CAB_PT, _BRANCO, _CENTRO
        c2.fill, c2.font, c2.alignment = _CAB_ZH, _BRANCO_ZH, _CENTRO


def _valor_celula(chave: str, valor: Any) -> Any:
    if valor is None or (isinstance(valor, float) and valor != valor):
        return None
    if chave in _COLUNAS_DATA:
        dt = para_data(valor)
        return dt if dt else valor
    if chave == "data":
        dt = para_data(valor)
        return dt.date() if dt else valor
    return valor


def aba_de_dados(wb: Workbook, titulo: str, df: pd.DataFrame, chaves: list[str],
                 miniaturas: dict[str, tuple[bytes, int, int]] | None = None) -> Worksheet:
    """Aba com uma linha por registro, cabeçalho J&T em duas linhas, filtro, painel congelado e as miniaturas."""
    ws = wb.create_sheet(titulo)
    chaves = [c for c in chaves if c in df.columns]
    _cabecalho(ws, 1, 1, chaves)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    miniaturas = miniaturas or {}
    larguras = {c: max(len(config.COLUNAS.get(c, (c,))[0]) + 2, 12) for c in chaves}
    pares_foto = {"imagem": "link_imagem", "imagem_avaria": "link_imagem_avaria"}
    for i, registro in enumerate(df[chaves].itertuples(index=False), start=3):
        valores = dict(zip(chaves, registro))
        tem_foto = False
        for j, chave in enumerate(chaves, start=1):
            v = valores[chave]
            celula = ws.cell(row=i, column=j)
            celula.font, celula.border, celula.alignment = _NORMAL, _BORDA, _MEIO
            if chave in ("link_imagem", "link_imagem_avaria"):
                url = str(v or "")
                if url.startswith("http"):
                    celula.value = "Abrir imagem"
                    celula.hyperlink = url
                    celula.font = _LINK
                elif url.startswith("data:"):
                    celula.value = "Imagem na coluna ao lado"
                else:
                    celula.value = ""
                continue
            celula.value = _valor_celula(chave, v)
            if isinstance(celula.value, datetime):
                celula.number_format = "yyyy-mm-dd hh:mm:ss"
            elif chave == "data" and not isinstance(celula.value, str):
                celula.number_format = "yyyy-mm-dd"
            elif chave == "horas_ate_avaria" and isinstance(celula.value, float):
                celula.number_format = "0.00"
            if chave in pares_foto:
                url = str(valores.get(pares_foto[chave]) or "")
                if url.startswith("http"):
                    celula.hyperlink = url
                    celula.font = _LINK
                mini = miniaturas.get(url)
                if mini:
                    _ancorar_imagem(ws, mini[0], mini[1], mini[2], i, j, _LARGURAS[chave])
                    celula.alignment = Alignment(vertical="bottom", horizontal="center")
                    tem_foto = True
            if chave not in _LARGURAS and celula.value is not None:
                larguras[chave] = max(larguras[chave], min(len(str(celula.value)) + 2, 60))
        if tem_foto:
            ws.row_dimensions[i].height = _ALTURA_LINHA_FOTO
    for j, chave in enumerate(chaves, start=1):
        ws.column_dimensions[get_column_letter(j)].width = _LARGURAS.get(chave, larguras[chave])
    ultima = max(2, len(df) + 2)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(chaves))}{ultima}"
    ws.freeze_panes = "B3"
    ws.sheet_view.zoomScale = 90
    return ws


def _tabela_contagem(ws: Worksheet, linha: int, coluna: int, titulo: str, df: pd.DataFrame, rotulo_valor: str = "valor") -> tuple[int, int]:
    """Tabela (Valor · Quantidade · %) com o título em cima. Devolve (primeira, última) linha dos dados."""
    t = ws.cell(row=linha, column=coluna, value=titulo)
    t.font = Font(name=_FONTE, bold=True, size=11, color=VERMELHO)
    _cabecalho(ws, linha + 1, coluna, [rotulo_valor, "quantidade", "percentual"])
    inicio = linha + 3
    for k, reg in enumerate(df.itertuples(index=False)):
        valores = (reg.valor, int(reg.quantidade), float(reg.percentual) / 100)
        for j, v in enumerate(valores):
            c = ws.cell(row=inicio + k, column=coluna + j, value=v)
            c.font, c.border = _NORMAL, _BORDA
            if j == 2:
                c.number_format = "0.0%"
    return inicio, inicio + max(len(df), 1) - 1


def _grafico(ws: Worksheet, titulo: str, coluna: int, primeira: int, ultima: int, ancora: str, *, cores: list[str] | None = None,
             horizontal: bool = False, rotulo_y: str = "Quantidade · 数量", coluna_valor: int | None = None,
             largura: float = 17.0, altura: float = 7.5, formato: str | None = None) -> None:
    grafico = BarChart()
    grafico.type = "bar" if horizontal else "col"
    grafico.title = titulo
    grafico.style = 10
    grafico.width, grafico.height = largura, altura
    dados = Reference(ws, min_col=coluna_valor or coluna + 1, min_row=primeira, max_row=ultima)
    categorias = Reference(ws, min_col=coluna, min_row=primeira, max_row=ultima)
    grafico.add_data(dados, titles_from_data=False)
    grafico.set_categories(categorias)
    serie = grafico.series[0]
    serie.graphicalProperties.solidFill = VERMELHO
    serie.graphicalProperties.line.solidFill = VERMELHO
    for i, cor in enumerate(cores or []):
        ponto = DataPoint(idx=i)
        ponto.graphicalProperties.solidFill = cor.lstrip("#")
        ponto.graphicalProperties.line.solidFill = cor.lstrip("#")
        serie.dPt.append(ponto)
    grafico.dataLabels = DataLabelList()
    grafico.dataLabels.showVal = True
    for campo in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(grafico.dataLabels, campo, False)
    if formato:
        grafico.dataLabels.numFmt = formato
    grafico.legend = None
    grafico.gapWidth = 60
    grafico.y_axis.title = rotulo_y
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    if horizontal:
        grafico.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(grafico, ancora)


class _Painel:
    """Aba "Resumo": blocos (tabela à esquerda, gráfico à direita) um embaixo do outro."""

    def __init__(self, ws: Worksheet):
        self.ws = ws
        self.linha = 4

    def secao(self, titulo: str) -> None:
        c = self.ws.cell(row=self.linha, column=1, value=titulo)
        c.font = Font(name=_FONTE, bold=True, size=13, color="FFFFFF")
        for j in range(1, 14):
            self.ws.cell(row=self.linha, column=j).fill = _CAB_PT
        self.ws.row_dimensions[self.linha].height = 22
        self.linha += 2

    def indicadores(self, titulo: str, pares: list[tuple[str, str, Any]]) -> None:
        t = self.ws.cell(row=self.linha, column=1, value=titulo)
        t.font = Font(name=_FONTE, bold=True, size=11, color=VERMELHO)
        _cabecalho(self.ws, self.linha + 1, 1, ["indicador", "valor"])
        for k, (pt, zh, valor) in enumerate(pares):
            c1 = self.ws.cell(row=self.linha + 3 + k, column=1, value=f"{pt} · {zh}")
            c2 = self.ws.cell(row=self.linha + 3 + k, column=2, value=valor)
            for c in (c1, c2):
                c.font, c.border = _NORMAL, _BORDA
            c2.font = Font(name=_FONTE, bold=True, size=11)
        self.linha += len(pares) + 5

    def bloco(self, titulo: str, df: pd.DataFrame, *, cores: list[str] | None = None, horizontal: bool = False,
              rotulo_valor: str = "valor") -> None:
        if df.empty:
            self.ws.cell(row=self.linha, column=1, value=f"{titulo}: sem dados").font = _NORMAL
            self.linha += 2
            return
        primeira, ultima = _tabela_contagem(self.ws, self.linha, 1, titulo, df, rotulo_valor)
        altura_cm = max(7.5, min(22.0, 1.2 + 0.55 * len(df))) if horizontal else 7.5
        _grafico(self.ws, titulo, 1, primeira, ultima, f"E{self.linha}", cores=cores, horizontal=horizontal, altura=altura_cm)
        linhas_grafico = int(altura_cm / 0.53) + 2
        self.linha += max(len(df) + 5, linhas_grafico)

    def bloco_valores(self, titulo: str, rotulos: list[str], valores: list[float | None], rotulo_y: str,
                      cores: list[str] | None = None, cab_valor: str = "mediana_h") -> None:
        t = self.ws.cell(row=self.linha, column=1, value=titulo)
        t.font = Font(name=_FONTE, bold=True, size=11, color=VERMELHO)
        _cabecalho(self.ws, self.linha + 1, 1, ["valor", cab_valor])
        inicio = self.linha + 3
        for k, (r, v) in enumerate(zip(rotulos, valores)):
            c1 = self.ws.cell(row=inicio + k, column=1, value=r)
            c2 = self.ws.cell(row=inicio + k, column=2, value=v if v is not None else 0)
            c2.number_format = "0.00"
            for c in (c1, c2):
                c.font, c.border = _NORMAL, _BORDA
        _grafico(self.ws, titulo, 1, inicio, inicio + len(rotulos) - 1, f"E{self.linha}", cores=cores, rotulo_y=rotulo_y,
                 formato="0.00")
        self.linha += max(len(rotulos) + 5, 16)


# ---------------------------------------------------------------------- planilha completa
def _cores_turno(df: pd.DataFrame) -> list[str]:
    return [config.CORES_TURNO.get(v, config.COR_NEUTRA) for v in df["valor"]]


def _cores_status(df: pd.DataFrame, mapa: dict[str, str]) -> list[str]:
    return [mapa.get(v, config.COR_NEUTRA) for v in df["valor"]]


CORES_TEM = {config.SIM: config.COR_BOM, config.NAO: config.COR_CRITICO, config.SEM_RASTREIO: config.COR_NEUTRA,
             config.ERRO_CONSULTA: config.COR_ALERTA}
CORES_JANELA = {config.JANELA_DENTRO: config.COR_BOM, config.JANELA_FORA: config.COR_CRITICO,
                config.JANELA_ANTES: config.COR_ALERTA, config.JANELA_SEM_DADOS: config.COR_NEUTRA}


def gerar_planilha(*, pesquisa: pd.DataFrame | None = None, segmentos: pd.DataFrame | None = None,
                   avaria: pd.DataFrame | None = None, resumo_av: dict[str, Any] | None = None,
                   erros: pd.DataFrame | None = None, prefixos: tuple[str, ...] | list[str] = config.PREFIXOS_JUNTAR,
                   info: dict[str, Any] | None = None, miniaturas: dict[str, tuple[bytes, int, int]] | None = None) -> bytes:
    """Monta o .xlsx (bytes). Abas: Resumo (gráficos) e uma aba de dados por análise, mais os erros."""
    info = info or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo 汇总"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 11
    titulo = ws.cell(row=1, column=1, value="J&T Express · JMS Analisador · 分析报告")
    titulo.font = Font(name=_FONTE, bold=True, size=16, color=VERMELHO)
    sub = (f"Gerado em {info.get('gerado_em', datetime.now().strftime('%d/%m/%Y %H:%M:%S'))} · "
           f"{info.get('remessas', 0)} remessa(s) · Turnos: " + " · ".join(f"{t} {h}" for t, h in config.HORARIO_TURNOS.items()))
    ws.cell(row=2, column=1, value=sub).font = Font(name=_FONTE, size=9, color="52514E")
    # Impressão/PDF: a largura toda (tabelas e gráficos) numa página.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    painel = _Painel(ws)

    if pesquisa is not None and not pesquisa.empty:
        painel.secao("Pesquisa por texto no Rastreamento do pacote · 轨迹文本搜索")
        for padrao in list(dict.fromkeys(pesquisa["padrao"]))[:5]:
            parte = pesquisa[pesquisa["padrao"] == padrao]
            achou = parte[parte["tem_padrao"] == config.SIM]
            painel.indicadores(f"Texto: “{padrao}”", [
                ("Remessas pesquisadas", "查询运单", len(parte)),
                ("Com o texto", "包含", len(achou)),
                ("% com o texto", "包含占比", f"{len(achou) / len(parte) * 100:.1f}%" if len(parte) else "—"),
                ("Turno com mais 1º bipes", "最多班次", achou["turno"].value_counts().idxmax() if len(achou) else "—"),
            ])
            tem = contar(parte["tem_padrao"], ignorar=(), ordem=list(CORES_TEM))
            painel.bloco("Tem o texto? · 是否包含", tem, cores=_cores_status(tem, CORES_TEM))
            turnos = contar(achou["turno"], ignorar=(), ordem=[*config.TURNOS, config.SEM_TURNO])
            painel.bloco("Turno do 1º bipe · 首次扫描班次", turnos, cores=_cores_turno(turnos))
            painel.bloco("Hora do 1º bipe · 首次扫描时段", contar_por_hora(achou["hora"]))
            painel.bloco("Quem escaneou (top 15) · 扫描人", contar(achou["login"], topo=15), horizontal=True)
            painel.bloco("Origem (top 10) · 扫描网点", contar(achou["origem"], topo=10), horizontal=True)
            painel.bloco("Próxima parada (top 10) · 下一站", contar(achou["proxima_parada"], topo=10), horizontal=True)

    if segmentos is not None and not segmentos.empty:
        painel.secao("Informação básica · 3 Segmentos · 三段码")
        com = segmentos[segmentos["cod_seg1"] != config.SEM_INFO]
        juntos = segmentos_juntados(segmentos, prefixos)
        painel.indicadores("Resumo dos segmentos", [
            ("Remessas pesquisadas", "查询运单", len(segmentos)),
            ("Com código de segmento", "有三段码", len(com)),
            ("Sem código", "无三段码", len(segmentos) - len(com)),
            (f"Com prefixo {', '.join(prefixos)}", "合并段码", len(juntos)),
        ])
        painel.bloco("Código de primeiro segmento · 一段码", contar(segmentos["cod_seg1"], topo=30), horizontal=True)
        painel.bloco(f"1º + 2º segmento juntos ({', '.join(prefixos)}) · 一段码+二段码", contar(juntos["seg_combinado"], topo=30),
                     horizontal=True)
        painel.bloco("SC Destino · 目的分拨中心", contar(segmentos["sc_destino"], topo=15), horizontal=True)
        painel.bloco("Base Destino (top 15) · 目的网点", contar(segmentos["base_destino"], topo=15), horizontal=True)

    if avaria is not None and not avaria.empty:
        r = resumo_av or {}
        jan = r.get("janela_h", config.JANELA_AVARIA_H)
        painel.secao(f"Avaria: tempo até o bipe de avaria (prazo {jan:g} h) · 破损扫描时效")
        fmt = lambda v, s="": "—" if v is None else f"{v:.2f}{s}".replace(".", ",")  # noqa: E731
        painel.indicadores("Resumo da avaria", [
            ("Remessas pesquisadas", "查询运单", r.get("remessas", len(avaria))),
            ("Com bipe de avaria", "有破损扫描", r.get("com_avaria", 0)),
            ("Medidas (início e avaria)", "可计算", r.get("medidas", 0)),
            ("Mediana de horas até a avaria", "中位小时", fmt(r.get("mediana_h"), " h")),
            ("Média de horas", "平均小时", fmt(r.get("media_h"), " h")),
            (f"Dentro de {jan:g} h", "时限内", r.get("dentro", 0)),
            (f"% dentro de {jan:g} h", "时限内占比", "—" if r.get("pct_dentro") is None else f"{r['pct_dentro']:.1f}%"),
        ])
        jan_df = contar(avaria["na_janela"], ignorar=(), ordem=list(CORES_JANELA))
        painel.bloco("Dentro do prazo? · 是否在时限内", jan_df, cores=_cores_status(jan_df, CORES_JANELA))
        med = r.get("mediana_por_turno") or {}
        painel.bloco_valores("Mediana de horas até a avaria por turno · 各班次中位小时", list(config.TURNOS),
                             [med.get(t) for t in config.TURNOS], "Horas · 小时",
                             cores=[config.CORES_TURNO[t] for t in config.TURNOS])
        turnos = contar(avaria["turno_avaria"], ignorar=(config.SEM_TURNO,), ordem=list(config.TURNOS))
        painel.bloco("Turno do bipe de avaria · 破损扫描班次", turnos, cores=_cores_turno(turnos))
        painel.bloco("Bipe usado como início · 起算扫描", contar(avaria["referencia"], ignorar=()))
        painel.bloco("Quem registrou a avaria (top 15) · 登记人", contar(avaria["login_avaria"], topo=15), horizontal=True)

    if pesquisa is not None and not pesquisa.empty:
        aba_de_dados(wb, "Pesquisa 搜索明细", pesquisa, list(pesquisa.columns), miniaturas)
    if segmentos is not None and not segmentos.empty:
        aba_de_dados(wb, "Segmentos 三段码", segmentos, config.COLUNAS_SEGMENTOS)
    if avaria is not None and not avaria.empty:
        aba_de_dados(wb, "Avaria 破损明细", avaria, config.COLUNAS_AVARIA, miniaturas)
    if erros is not None and not erros.empty:
        aba_de_dados(wb, "Erros 错误", erros, config.COLUNAS_ERROS)

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def urls_de_imagens(*dfs: pd.DataFrame | None) -> list[str]:
    urls: list[str] = []
    for df in dfs:
        if df is None or df.empty:
            continue
        for coluna in ("link_imagem", "link_imagem_avaria"):
            if coluna in df.columns:
                urls.extend(u for u in df[coluna].astype(str) if u.startswith(("http", "data:")))
    return list(dict.fromkeys(urls))
